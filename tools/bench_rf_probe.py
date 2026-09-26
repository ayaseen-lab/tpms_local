"""Program one row, then list every sensor ID the board reports over RF.

Shows whether an ID came from a fresh SRAM frame or a stale leftover.

Run:  python tools/bench_rf_probe.py <excel_row>
"""

from __future__ import annotations

import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [
    str(ROOT / "sdr_ui"),
    str(ROOT / "hamaton-sdk-python-fix-uart-transport-timing" / "src"),
    str(ROOT),
]

from openpyxl import load_workbook  # noqa: E402

from hamaton.commands.receive_rf import ReceiveRfCodec  # noqa: E402
from hamaton.commands.trigger import TriggerCodec  # noqa: E402
from hamaton.frame import HamatonFrame  # noqa: E402
from hamaton.models import ProgramResult  # noqa: E402
from tpms_bench import jlink_ram  # noqa: E402
from tpms_bench.board import BoardSession, generate_unique_oeid  # noqa: E402
from tpms_bench.excel_io import parse_code  # noqa: E402
from tpms_bench.hardware import preferred_usb_ttl_port  # noqa: E402
from tpms_bench.query_sensor import QuerySensorCodec  # noqa: E402
from tpms_bench.uart import open_uart  # noqa: E402

XLSX = Path(
    "/Users/mr.macbook/Downloads/hamaton-sdk-python-fix-uart-transport-timing"
    "/Hamaton_database_20260126_1305.xlsx"
)


def log(msg: str) -> None:
    print(msg, flush=True)


def row_codes(excel_row: int):
    wb = load_workbook(XLSX, read_only=True)
    ws = wb[wb.sheetnames[0]]
    header = {}
    for idx, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1)), start=1):
        header[str(cell.value or "").strip().upper().replace(" ", "")] = idx
    codes = tuple(
        parse_code(ws.cell(excel_row, header.get(n, d)).value)
        for n, d in (("CODEA", 16), ("CODEB", 17), ("CODEC", 18))
    )
    label = " ".join(
        str(ws.cell(excel_row, header.get(n, d)).value or "")
        for n, d in (("MAKE", 1), ("MODEL", 2))
    )
    wb.close()
    return codes, label


def all_board_frames(blob: bytes):
    return jlink_ram.extract_frames(blob, HamatonFrame.BOARD_ADDRESS)


def main() -> int:
    excel_row = int(sys.argv[1]) if len(sys.argv) > 1 else 4
    codes, label = row_codes(excel_row)
    log(f"row {excel_row}: {label} codes={' '.join(f'{c:08X}' for c in codes)}")

    ser = open_uart(preferred_usb_ttl_port())
    session = BoardSession(ser)
    session.jlink_enabled = True
    session.jlink_ok = True
    session.uart_rx_active = False
    try:
        result = session.program(*codes, generate_unique_oeid(set()))
        if not (result.is_success and isinstance(result.value, ProgramResult)):
            log(f"program failed: {result.message}")
            return 1
        bits = getattr(result.value, "id_bit_length", 0x20)
        log(
            f"programmed id={result.value.sensor_id.hex().upper()} "
            f"id_bits=0x{bits:02X}"
        )

        before = {
            (o, f.to_bytes()) for o, f in all_board_frames(jlink_ram.dump_sram_result().blob)
        }
        log(f"snapshot frames: {len(before)}")

        session.send_codec(ReceiveRfCodec(True, *codes))
        session.send_codec(TriggerCodec(*codes))

        trigger = TriggerCodec(*codes)
        query = QuerySensorCodec()
        seen: set[bytes] = set()
        deadline = time.monotonic() + 25.0
        while time.monotonic() < deadline:
            dump = jlink_ram.dump_sram_result()
            if dump.ok:
                for offset, frame in all_board_frames(dump.blob):
                    raw = frame.to_bytes()
                    if raw in seen:
                        continue
                    seen.add(raw)
                    fresh = (offset, raw) not in before
                    for name, codec in (("trigger", trigger), ("query", query)):
                        if not codec.accepts_response(frame):
                            continue
                        try:
                            parsed = codec.parse_response(frame)
                        except Exception as exc:
                            log(
                                f"  {'NEW  ' if fresh else 'stale'} @0x{offset:05X} "
                                f"{name} DROPPED {type(exc).__name__}: {exc} "
                                f"payload={frame.payload.hex()}"
                            )
                            continue
                        value = parsed.value
                        sid = getattr(value, "sensor_id", None)
                        log(
                            f"  {'NEW  ' if fresh else 'stale'} @0x{offset:05X} "
                            f"{name} state={parsed.state.value} "
                            f"id={sid.hex().upper() if sid else '—'} "
                            f"payload={frame.payload.hex()}"
                        )
            time.sleep(0.3)
        session.send_codec(ReceiveRfCodec(False, *codes))
    finally:
        try:
            session.idle_board(hard=False)
        except Exception:
            pass
        ser.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
