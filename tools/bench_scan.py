"""Scan Excel rows: does the board reply to Program One Sensor at all?

Run:  python tools/bench_scan.py [first_row] [last_row] [wait_s]
"""

from __future__ import annotations

import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [
    str(ROOT / "sdr_ui"),
    str(ROOT / "hamaton-sdk-python-fix-uart-transport-timing" / "src"),
    str(ROOT),
]

from openpyxl import load_workbook  # noqa: E402

from hamaton.commands.program_sensor import ProgramOneSensorCodec  # noqa: E402
from hamaton.frame import HamatonFrame  # noqa: E402
from tpms_bench import jlink_ram  # noqa: E402
from tpms_bench.board import BoardSession, generate_unique_oeid  # noqa: E402
from tpms_bench.excel_io import parse_code  # noqa: E402
from tpms_bench.hardware import preferred_usb_ttl_port  # noqa: E402
from tpms_bench.program_search import ProgramSearchCodec  # noqa: E402
from tpms_bench.uart import open_uart  # noqa: E402

XLSX = Path(
    "/Users/mr.macbook/Downloads/hamaton-sdk-python-fix-uart-transport-timing"
    "/Hamaton_database_20260126_1305.xlsx"
)


def log(msg: str) -> None:
    print(msg, flush=True)


def load_rows(first: int, last: int) -> list[dict]:
    wb = load_workbook(XLSX, read_only=True)
    ws = wb[wb.sheetnames[0]]
    header = {}
    for idx, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1)), start=1):
        header[str(cell.value or "").strip().upper().replace(" ", "")] = idx
    rows = []
    for excel_row in range(first, last + 1):
        codes = tuple(
            parse_code(ws.cell(excel_row, header.get(n, d)).value)
            for n, d in (("CODEA", 16), ("CODEB", 17), ("CODEC", 18))
        )
        rows.append(
            {
                "row": excel_row,
                "label": " ".join(
                    str(ws.cell(excel_row, header.get(n, d)).value or "")
                    for n, d in (("MAKE", 1), ("MODEL", 2))
                )[:22],
                "freq": str(ws.cell(excel_row, header.get("FREQ", 10)).value or ""),
                "codes": codes,
            }
        )
    wb.close()
    return rows


def frames_for(blob: bytes, wanted: tuple[int, ...]) -> list[tuple[int, HamatonFrame]]:
    out = []
    for offset, frame in jlink_ram.extract_frames(blob, HamatonFrame.BOARD_ADDRESS):
        payload = frame.payload
        if len(payload) >= 2 and payload[0] in wanted:
            out.append((offset, frame))
    return out


def wait_reply(session, codec, wanted, seconds: float):
    """Poll SRAM until a new frame for this command shows up."""
    before = {
        (o, f.to_bytes()) for o, f in frames_for(jlink_ram.dump_sram_result().blob, wanted)
    }
    session.send_codec(codec)
    deadline = time.monotonic() + seconds
    while time.monotonic() < deadline:
        dump = jlink_ram.dump_sram_result()
        if dump.ok:
            for offset, frame in reversed(frames_for(dump.blob, wanted)):
                if (offset, frame.to_bytes()) in before:
                    continue
                try:
                    return codec.parse_response(frame), frame
                except Exception as exc:
                    return f"DROPPED {type(exc).__name__}: {exc}", frame
        time.sleep(0.15)
    return None, None


def main() -> int:
    first = int(sys.argv[1]) if len(sys.argv) > 1 else 2
    last = int(sys.argv[2]) if len(sys.argv) > 2 else 13
    wait_s = float(sys.argv[3]) if len(sys.argv) > 3 else 8.0

    rows = load_rows(first, last)
    ser = open_uart(preferred_usb_ttl_port())
    session = BoardSession(ser)
    session.jlink_enabled = True
    session.jlink_ok = True
    session.uart_rx_active = False

    replied = 0
    silent: list[int] = []
    try:
        for entry in rows:
            if None in entry["codes"]:
                log(f"row {entry['row']:>3} {entry['label']:<22} SKIP (missing codes)")
                continue
            session.idle_board(hard=False)
            oeid = generate_unique_oeid(set())
            codec = ProgramOneSensorCodec(*entry["codes"], oeid)
            started = time.monotonic()
            result, frame = wait_reply(session, codec, (0x21, 0xA1), wait_s)
            took = time.monotonic() - started
            codes_txt = " ".join(f"{c:08X}" for c in entry["codes"])
            if result is None:
                log(
                    f"row {entry['row']:>3} {entry['label']:<22} "
                    f"freq={entry['freq'] or '—':<8} SILENT  {took:4.1f}s  {codes_txt}"
                )
                silent.append(entry["row"])
                # Does the board answer the LF search variant instead?
                search = ProgramSearchCodec(*entry["codes"])
                sres, _ = wait_reply(session, search, (0x21, 0xA1), 6.0)
                log(f"        search reply: {sres if sres is None else getattr(sres, 'state', sres)}")
            else:
                state = getattr(result, "state", None)
                txt = state.value if state else str(result)
                log(
                    f"row {entry['row']:>3} {entry['label']:<22} "
                    f"freq={entry['freq'] or '—':<8} REPLY {txt:<8} {took:4.1f}s  "
                    f"payload={frame.payload.hex()}"
                )
                replied += 1
    finally:
        try:
            session.idle_board(hard=False)
        except Exception:
            pass
        ser.close()

    log(f"\nreplied {replied}/{len(rows)} · silent rows: {silent}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
