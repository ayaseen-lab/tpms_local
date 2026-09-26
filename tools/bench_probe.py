"""Dump the raw board replies for one Excel row so dropped frames become visible.

Run:  python tools/bench_probe.py <excel_row> [seconds]
"""

from __future__ import annotations

import sys
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [
    str(ROOT / "sdr_ui"),
    str(ROOT / "hamaton-sdk-python-fix-uart-transport-timing" / "src"),
    str(ROOT),
]

from openpyxl import load_workbook  # noqa: E402

from hamaton.commands.program_sensor import ProgramOneSensorCodec  # noqa: E402
from hamaton.frame import HamatonFrame  # noqa: E402
from tpms_bench import jlink_ram  # noqa: E402
from tpms_bench.board import BoardSession, generate_unique_oeid  # noqa: E402
from tpms_bench.excel_io import parse_code  # noqa: E402
from tpms_bench.hardware import preferred_usb_ttl_port  # noqa: E402
from tpms_bench.uart import open_uart  # noqa: E402

XLSX = Path(
    "/Users/mr.macbook/Downloads/hamaton-sdk-python-fix-uart-transport-timing"
    "/Hamaton_database_20260126_1305.xlsx"
)


def log(msg: str) -> None:
    print(msg, flush=True)


def row_codes(excel_row: int):
    wb = load_workbook(XLSX, read_only=True)
    ws = wb[wb.sheetnames[0]]
    header = {}
    for idx, cell in enumerate(next(ws.iter_rows(min_row=1, max_row=1)), start=1):
        header[str(cell.value or "").strip().upper().replace(" ", "")] = idx
    codes = tuple(
        parse_code(ws.cell(excel_row, header.get(name, default)).value)
        for name, default in (("CODEA", 16), ("CODEB", 17), ("CODEC", 18))
    )
    label = " ".join(
        str(ws.cell(excel_row, header.get(n, d)).value or "")
        for n, d in (("MAKE", 1), ("MODEL", 2))
    )
    freq = ws.cell(excel_row, header.get("FREQ", 10)).value
    wb.close()
    return codes, label, freq


def program_frames(blob: bytes) -> list[tuple[int, HamatonFrame]]:
    out = []
    for offset, frame in jlink_ram.extract_frames(blob, HamatonFrame.BOARD_ADDRESS):
        payload = frame.payload
        if len(payload) >= 2 and payload[0] in (0x21, 0xA1):
            out.append((offset, frame))
    return out


def main() -> int:
    excel_row = int(sys.argv[1]) if len(sys.argv) > 1 else 3
    seconds = float(sys.argv[2]) if len(sys.argv) > 2 else 45.0

    codes, label, freq = row_codes(excel_row)
    log(f"row {excel_row}: {label} freq={freq}")
    log("codes: " + " ".join(f"{c:08X}" if c is not None else "None" for c in codes))
    if None in codes:
        log("row has no usable codes")
        return 1

    ser = open_uart(preferred_usb_ttl_port())
    session = BoardSession(ser)
    session.jlink_enabled = True
    session.jlink_ok = True
    session.uart_rx_active = False
    try:
        session.idle_board(hard=False)
        before = {
            (offset, frame.to_bytes()) for offset, frame in program_frames(
                jlink_ram.dump_sram_result().blob
            )
        }
        log(f"pre-existing 0x21 frames: {len(before)}")

        oeid = generate_unique_oeid(set())
        codec = ProgramOneSensorCodec(*codes, oeid)
        log(f"TX program oeid={oeid:08X}")
        session.send_codec(codec)

        deadline = time.monotonic() + seconds
        seen: set[bytes] = set()
        while time.monotonic() < deadline:
            uart = session._drain_uart(0.2)
            if uart:
                log(f"  UART bytes: {uart.hex()}")
            dump = jlink_ram.dump_sram_result()
            if dump.ok:
                for offset, frame in program_frames(dump.blob):
                    raw = frame.to_bytes()
                    if raw in seen:
                        continue
                    seen.add(raw)
                    fresh = (offset, raw) not in before
                    note = ""
                    try:
                        result = codec.parse_response(frame)
                        note = f"parsed state={result.state.value} msg={result.message}"
                    except Exception as exc:
                        note = f"DROPPED {type(exc).__name__}: {exc}"
                    log(
                        f"  [{time.monotonic() - (deadline - seconds):5.1f}s] "
                        f"{'NEW  ' if fresh else 'stale'} @0x{offset:05X} "
                        f"payload={frame.payload.hex()} · {note}"
                    )
            time.sleep(0.2)
    finally:
        try:
            session.idle_board(hard=False)
        except Exception:
            pass
        ser.close()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
