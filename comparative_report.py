"""Comparative Analysis — every Board row and every SDR ID with OK/NOK per source.

Current session: one table row for each TPMS Board result row, plus SDR-only IDs
that the Board never saw. Does not collapse multiple Board vehicles that share
(or lack) a Sensor ID.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from tpms_bench.compare import ids_related, norm_id

REPORT_KIND = "Comparative Analysis Report"
REPORT_SUBTITLE = (
    "Every Board row + every SDR ID — TPMS Board OK/NOK vs SDR OK/NOK"
)
REPORT_DISCLAIMER = (
    "The Comparison table lists every TPMS Board result from this session "
    "(one line per Board row), with the matching SDR OK/NOK when that Sensor ID "
    "was also heard by the SDR Receiver. Extra lines are added for SDR-only IDs "
    "the Board never reported. AGREE/DISAGREE apply when both sides have a result. "
    "Board time is seconds spent on that Bench row; SDR time is seconds from the "
    "first packet for that ID until OK (or until the latest packet if still NOK)."
)

HEADER_FILL = PatternFill("solid", fgColor="101011")
HEADER_FONT = Font(bold=True, color="FFFFFF")
OK_FILL = PatternFill("solid", fgColor="ECFDF5")
LOW_FILL = PatternFill("solid", fgColor="FEF2F2")
WARN_FILL = PatternFill("solid", fgColor="FFF7ED")
NEUTRAL_FILL = PatternFill("solid", fgColor="F8FAFC")

NAVY = colors.HexColor("#101011")
ORANGE = colors.HexColor("#00D3BF")
LIGHT = colors.HexColor("#EEF1F5")
OK_BG = colors.HexColor("#E6F6EC")
FAIL_BG = colors.HexColor("#FEE2E2")
LINE = colors.HexColor("#D0D7E0")

COMPARISON_HEADERS = [
    "Board #",
    "Vehicle / Protocol",
    "Sensor ID",
    "Seen by",
    "TPMS Board result",
    "TPMS Board reason",
    "SDR result",
    "SDR reason",
    "rtl_433 Decoder",
    "Board time (s)",
    "SDR time (s)",
    "Comparison",
]


@dataclass
class CompareRow:
    sensor_id: str
    seen_by: str = ""
    board_result: str = ""
    board_reason: str = ""
    sdr_result: str = ""
    sdr_reason: str = ""
    verdict: str = ""
    excel_row: int = 0
    vehicle: str = ""
    sdr_decoder: str = ""
    board_time_s: float | None = None
    sdr_time_s: float | None = None


def _format_seconds(seconds: float | None) -> str:
    if seconds is None:
        return "—"
    value = max(0.0, float(seconds))
    if value < 60:
        return f"{value:.1f}"
    mins, rem = divmod(int(round(value)), 60)
    return f"{mins}m {rem:02d}s"


def _parse_duration(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().lower().replace("s", "")
    if text in {"na", "n/a", "—", "-"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _sdr_acquire_seconds(snapshot: Any) -> float | None:
    if hasattr(snapshot, "acquire_seconds"):
        return _parse_duration(getattr(snapshot, "acquire_seconds", None))
    if isinstance(snapshot, dict):
        return _parse_duration(
            snapshot.get("acquire_seconds")
            or snapshot.get("duration_s")
            or snapshot.get("time_s")
        )
    return None


@dataclass
class ComparisonResult:
    rows: list[CompareRow] = field(default_factory=list)
    agree: int = 0
    disagree: int = 0
    board_only: int = 0
    sdr_only: int = 0
    board_rows: int = 0
    sdr_ids: int = 0
    # Board rows that had an SDR side (agree + disagree) — not unique sensor IDs.
    sdr_matches: int = 0
    generated_at: datetime = field(default_factory=datetime.now)

    @property
    def matched(self) -> int:
        return self.agree + self.disagree

    @property
    def agreement_rate(self) -> float:
        if self.matched <= 0:
            return 0.0
        return 100.0 * self.agree / self.matched


def _sdr_side(snapshot: Any) -> tuple[str, str]:
    reading = snapshot
    if hasattr(reading, "qualifies_ok"):
        ok = reading.qualifies_ok()
        reason = reading.nok_reason() if hasattr(reading, "nok_reason") else ""
        return ("OK" if ok else "NOK"), reason or ""
    if isinstance(snapshot, dict):
        result = str(snapshot.get("result") or snapshot.get("ok_nok") or "").upper()
        reason = str(snapshot.get("reason") or snapshot.get("nok_reason") or "")
        if result in ("OK", "NOK"):
            return result, reason
        return "NOK", reason or "incomplete"
    return "NOK", "incomplete"


def _sdr_protocol(snapshot: Any) -> str:
    """Full rtl_433 library decoder label for this SDR reading."""
    if hasattr(snapshot, "display_decoder"):
        label = str(getattr(snapshot, "display_decoder", "") or "").strip()
        if label and label not in {"—", "-"}:
            return label
    if hasattr(snapshot, "decoder"):
        label = str(getattr(snapshot, "decoder", "") or "").strip()
        if label and label not in {"—", "-"}:
            return label
    if hasattr(snapshot, "model"):
        model = str(getattr(snapshot, "model", None) or "").strip()
        protocol_id = getattr(snapshot, "protocol_id", None)
        try:
            from config import format_rtl433_decoder

            return format_rtl433_decoder(protocol_id, model)
        except Exception:
            return model
    if isinstance(snapshot, dict):
        for key in ("decoder", "display_decoder", "rtl_433_decoder"):
            label = str(snapshot.get(key) or "").strip()
            if label and label not in {"—", "-"}:
                return label
        model = str(snapshot.get("model") or snapshot.get("protocol") or "").strip()
        protocol_id = snapshot.get("protocol_id")
        try:
            from config import format_rtl433_decoder

            return format_rtl433_decoder(
                int(protocol_id) if protocol_id not in (None, "") else None,
                model,
            )
        except Exception:
            return model
    return ""


def _tally(result: ComparisonResult, verdict: str) -> None:
    if verdict == "AGREE":
        result.agree += 1
    elif verdict == "DISAGREE":
        result.disagree += 1
    elif verdict == "BOARD_ONLY":
        result.board_only += 1
    elif verdict == "SDR_ONLY":
        result.sdr_only += 1


def build_comparison(
    sdr_sensors: dict[str, Any] | Iterable[Any],
    board_rows: list[dict],
) -> ComparisonResult:
    """One line per Board session row, plus SDR-only IDs not on any Board row."""
    # display, result, reason, protocol, acquire_seconds
    sdr_map: dict[str, tuple[str, str, str, str, float | None]] = {}
    if isinstance(sdr_sensors, dict):
        items = list(sdr_sensors.items())
    else:
        items = []
        for reading in sdr_sensors:
            sid = getattr(reading, "sensor_id", None) or (
                reading.get("sensor_id") if isinstance(reading, dict) else ""
            )
            items.append((sid, reading))

    for sid, reading in items:
        key = norm_id(str(sid or ""))
        if not key:
            continue
        protocol = _sdr_protocol(reading)
        # Ignore Board-mirrored placeholders — only real rtl_433 library rows.
        if "waiting for decode" in protocol.lower() or protocol.lower().startswith("board rf"):
            continue
        if not protocol.startswith("["):
            # Still allow typed OK/NOK snapshots without a library tag.
            model = str(getattr(reading, "model", "") or "")
            if not model or model.upper() == "TPMS":
                continue
        result, reason = _sdr_side(reading)
        display = str(getattr(reading, "sensor_id", None) or sid or key)
        acquire = _sdr_acquire_seconds(reading)
        sdr_map[key] = (display, result, reason, protocol, acquire)

    out = ComparisonResult()
    out.board_rows = len(board_rows)
    out.sdr_ids = len(sdr_map)
    board_seen_keys: set[str] = set()

    # Preserve Board session order (as tested).
    for brow in board_rows:
        sid = str(brow.get("sensor_id") or "").strip()
        key = norm_id(sid)
        board_res = str(brow.get("board_performance") or "").upper()
        board_reason = str(brow.get("nok_reason") or "")
        excel_row = int(brow.get("excel_row") or 0)
        vehicle = f"{brow.get('make') or ''} {brow.get('model') or ''}".strip() or "—"
        display = sid if sid else "(no ID)"
        board_time = _parse_duration(brow.get("duration_s"))

        sdr_res = ""
        sdr_reason = ""
        sdr_time = None
        sdr_decoder = ""
        sdr_key = ""
        if key:
            board_seen_keys.add(key)
            hit = sdr_map.get(key)
            if hit is None:
                for cand_key, cand in sdr_map.items():
                    if ids_related(sid, cand[0]) or ids_related(sid, cand_key):
                        hit = cand
                        sdr_key = cand_key
                        break
            else:
                sdr_key = key
            if hit is not None:
                display_sdr, sdr_res, sdr_reason, sdr_decoder, sdr_time = hit
                display = display or display_sdr
                board_seen_keys.add(sdr_key)

        if key and sdr_res:
            seen_by = "Both"
            verdict = "AGREE" if board_res == sdr_res else "DISAGREE"
        else:
            seen_by = "TPMS Board only"
            verdict = "BOARD_ONLY"

        _tally(out, verdict)
        out.rows.append(
            CompareRow(
                sensor_id=display,
                seen_by=seen_by,
                board_result=board_res,
                board_reason=board_reason,
                sdr_result=sdr_res,
                sdr_reason=sdr_reason,
                verdict=verdict,
                excel_row=excel_row,
                vehicle=vehicle,
                sdr_decoder=sdr_decoder or "—",
                board_time_s=board_time,
                sdr_time_s=sdr_time,
            )
        )

    # SDR IDs never reported on any Board row this session.
    for key, (display, sdr_res, sdr_reason, protocol, sdr_time) in sorted(
        sdr_map.items(), key=lambda x: x[0]
    ):
        if key in board_seen_keys:
            continue
        _tally(out, "SDR_ONLY")
        out.rows.append(
            CompareRow(
                sensor_id=display,
                seen_by="SDR only",
                board_result="",
                board_reason="",
                sdr_result=sdr_res,
                sdr_reason=sdr_reason,
                verdict="SDR_ONLY",
                excel_row=0,
                vehicle=protocol or "—",
                sdr_decoder=protocol or "—",
                board_time_s=None,
                sdr_time_s=sdr_time,
            )
        )

    out.sdr_matches = out.agree + out.disagree
    return out


def _style_header(ws, headers: list[str]) -> None:
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)


def _set_widths(ws, widths: dict[int, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def _fill_for_verdict(verdict: str) -> PatternFill | None:
    if verdict == "AGREE":
        return OK_FILL
    if verdict == "DISAGREE":
        return LOW_FILL
    if verdict in ("BOARD_ONLY", "SDR_ONLY"):
        return WARN_FILL
    return NEUTRAL_FILL


def export_comparative_excel(path: str | Path, comparison: ComparisonResult) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    cmp_ws = wb.active
    cmp_ws.title = "Comparison"
    _style_header(cmp_ws, COMPARISON_HEADERS)
    for row in comparison.rows:
        values = [
            row.excel_row or "—",
            row.vehicle or "—",
            row.sensor_id,
            row.seen_by,
            row.board_result or "—",
            row.board_reason,
            row.sdr_result or "—",
            row.sdr_reason,
            row.sdr_decoder or "—",
            _format_seconds(row.board_time_s),
            _format_seconds(row.sdr_time_s),
            row.verdict,
        ]
        cmp_ws.append(values)
        fill = _fill_for_verdict(row.verdict)
        if fill:
            for col in range(1, len(values) + 1):
                cmp_ws.cell(row=cmp_ws.max_row, column=col).fill = fill
    _set_widths(
        cmp_ws,
        {
            1: 10, 2: 22, 3: 16, 4: 14, 5: 14, 6: 22, 7: 12, 8: 22,
            9: 34, 10: 12, 11: 12, 12: 12,
        },
    )
    cmp_ws.freeze_panes = "A2"
    cmp_ws.auto_filter.ref = f"A1:L{max(cmp_ws.max_row, 1)}"

    info = wb.create_sheet("Report Info")
    info["A1"] = REPORT_KIND
    info["A1"].font = Font(bold=True, size=14)
    info["A2"] = REPORT_SUBTITLE
    info["A2"].font = Font(italic=True, size=10, color="64748B")
    info["A3"] = REPORT_DISCLAIMER
    info["A3"].font = Font(size=9, color="475569")
    info.merge_cells("A3:B3")
    rows_info = [
        ("Report type", "All Board rows + SDR-only IDs — OK/NOK per source"),
        ("Generated", comparison.generated_at.strftime("%Y-%m-%d %H:%M:%S")),
        ("Board rows in session", comparison.board_rows),
        ("SDR unique IDs in session", comparison.sdr_ids),
        ("Table lines", len(comparison.rows)),
        ("Agree", comparison.agree),
        ("Disagree", comparison.disagree),
        ("TPMS Board only", comparison.board_only),
        ("SDR only", comparison.sdr_only),
        ("Agreement rate (matched lines)", f"{comparison.agreement_rate:.1f}%"),
        (
            "Time columns",
            "Board time = seconds for that Bench row; SDR time = seconds from first packet to OK",
        ),
    ]
    for i, (key, val) in enumerate(rows_info, start=5):
        info[f"A{i}"] = key
        info[f"B{i}"] = val
    info.column_dimensions["A"].width = 34
    info.column_dimensions["B"].width = 56

    summary = wb.create_sheet("Summary")
    _style_header(summary, ["Metric", "Count"])
    for label, count in (
        ("Board rows in session", comparison.board_rows),
        ("SDR unique IDs in session", comparison.sdr_ids),
        ("Table lines (Board rows + SDR-only)", len(comparison.rows)),
        ("Agree", comparison.agree),
        ("Disagree", comparison.disagree),
        ("TPMS Board only", comparison.board_only),
        ("SDR only", comparison.sdr_only),
        ("Agreement rate %", round(comparison.agreement_rate, 1)),
    ):
        summary.append([label, count])
    _set_widths(summary, {1: 40, 2: 14})

    wb.save(path)
    return path


def _pdf_styles():
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "cmp_title",
            parent=base["Title"],
            fontName="Helvetica-Bold",
            fontSize=18,
            textColor=NAVY,
            alignment=TA_CENTER,
            spaceAfter=4,
        ),
        "sub": ParagraphStyle(
            "cmp_sub",
            parent=base["Normal"],
            fontSize=10,
            textColor=ORANGE,
            alignment=TA_CENTER,
            spaceAfter=8,
        ),
        "body": ParagraphStyle(
            "cmp_body",
            parent=base["Normal"],
            fontSize=8.5,
            textColor=NAVY,
            leading=11,
            alignment=TA_LEFT,
            spaceAfter=8,
        ),
        "h": ParagraphStyle(
            "cmp_h",
            parent=base["Heading2"],
            textColor=NAVY,
            fontSize=12,
            spaceBefore=8,
            spaceAfter=6,
        ),
        "tiny": ParagraphStyle(
            "cmp_tiny",
            parent=base["Normal"],
            fontSize=7,
            textColor=NAVY,
            leading=9,
        ),
        "th": ParagraphStyle(
            "cmp_th",
            parent=base["Normal"],
            fontName="Helvetica-Bold",
            fontSize=7,
            textColor=colors.white,
            leading=9,
            alignment=TA_CENTER,
        ),
    }


def _header_footer(canvas, doc):
    canvas.saveState()
    canvas.setFillColor(NAVY)
    canvas.rect(0, landscape(A4)[1] - 14 * mm, landscape(A4)[0], 14 * mm, fill=1, stroke=0)
    canvas.setFillColor(ORANGE)
    canvas.rect(0, landscape(A4)[1] - 16 * mm, landscape(A4)[0], 2 * mm, fill=1, stroke=0)
    canvas.setFillColor(colors.white)
    canvas.setFont("Helvetica-Bold", 10)
    canvas.drawString(
        14 * mm,
        landscape(A4)[1] - 9 * mm,
        "Comparative Analysis  ·  All Board rows + SDR IDs (OK/NOK)",
    )
    canvas.setFont("Helvetica", 8)
    canvas.drawRightString(
        landscape(A4)[0] - 14 * mm,
        landscape(A4)[1] - 9 * mm,
        datetime.now().strftime("%Y-%m-%d %H:%M"),
    )
    canvas.setFillColor(NAVY)
    canvas.rect(0, 0, landscape(A4)[0], 10 * mm, fill=1, stroke=0)
    canvas.setFillColor(colors.white)
    canvas.setFont("Helvetica", 8)
    canvas.drawString(14 * mm, 4 * mm, "Current session  ·  Board rows are not collapsed by ID")
    canvas.drawRightString(landscape(A4)[0] - 14 * mm, 4 * mm, f"Page {doc.page}")
    canvas.restoreState()


def export_comparative_pdf(path: str | Path, comparison: ComparisonResult) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    styles = _pdf_styles()
    doc = SimpleDocTemplate(
        str(path),
        pagesize=landscape(A4),
        leftMargin=8 * mm,
        rightMargin=8 * mm,
        topMargin=22 * mm,
        bottomMargin=16 * mm,
    )
    total_lines = len(comparison.rows)
    story = [
        Paragraph(REPORT_KIND, styles["title"]),
        Paragraph(REPORT_SUBTITLE, styles["sub"]),
        Paragraph(
            f"{REPORT_DISCLAIMER} "
            f"<b>Board rows: {comparison.board_rows} · SDR IDs: {comparison.sdr_ids} · "
            f"Table lines: {total_lines}.</b>",
            styles["body"],
        ),
    ]

    cards = [
        ["Board rows", str(comparison.board_rows)],
        ["SDR IDs", str(comparison.sdr_ids)],
        ["Table lines", str(total_lines)],
        ["Agree", str(comparison.agree)],
        ["Disagree", str(comparison.disagree)],
    ]
    card_table = Table(cards, colWidths=[52 * mm] * 5)
    card_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), LIGHT),
                ("BACKGROUND", (0, 1), (-1, 1), colors.white),
                ("TEXTCOLOR", (0, 0), (-1, -1), NAVY),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica"),
                ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("BOX", (0, 0), (-1, -1), 0.5, LINE),
                ("INNERGRID", (0, 0), (-1, -1), 0.4, LINE),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.append(card_table)
    story.append(Spacer(1, 6))
    story.append(
        Paragraph(
            f"Comparison table — all Board rows + SDR-only IDs ({total_lines} lines)",
            styles["h"],
        )
    )

    def cell(text: str, limit: int = 60) -> Paragraph:
        raw = str(text or "—").replace("\n", " ").strip()
        if len(raw) > limit:
            raw = raw[: limit - 1] + "…"
        return Paragraph(
            raw.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"),
            styles["tiny"],
        )

    header = [Paragraph(str(h), styles["th"]) for h in COMPARISON_HEADERS]
    body_rows: list[list] = []
    for row in comparison.rows:
        body_rows.append(
            [
                cell(row.excel_row or "—", 8),
                cell(row.vehicle, 22),
                cell(row.sensor_id, 16),
                cell(row.seen_by, 14),
                cell(row.board_result, 8),
                cell(row.board_reason, 24),
                cell(row.sdr_result, 8),
                cell(row.sdr_reason, 24),
                cell(row.sdr_decoder or "—", 34),
                cell(_format_seconds(row.board_time_s), 10),
                cell(_format_seconds(row.sdr_time_s), 10),
                cell(row.verdict, 12),
            ]
        )

    col_widths = [
        11 * mm,
        24 * mm,
        18 * mm,
        16 * mm,
        14 * mm,
        24 * mm,
        12 * mm,
        24 * mm,
        34 * mm,
        14 * mm,
        14 * mm,
        16 * mm,
    ]
    chunk_size = 16
    header_style = [
        ("BACKGROUND", (0, 0), (-1, 0), NAVY),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 7),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BOX", (0, 0), (-1, -1), 0.5, LINE),
        ("INNERGRID", (0, 0), (-1, -1), 0.3, LINE),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 2),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2),
    ]
    if not body_rows:
        empty = Table([header, [cell("—")] * len(COMPARISON_HEADERS)], colWidths=col_widths, repeatRows=1)
        empty.setStyle(TableStyle(header_style))
        story.append(empty)
    else:
        for start in range(0, len(body_rows), chunk_size):
            chunk = body_rows[start : start + chunk_size]
            data = [header] + chunk
            table = Table(data, colWidths=col_widths, repeatRows=1)
            style_cmds = list(header_style) + [
                ("FONTSIZE", (0, 1), (-1, -1), 7),
                ("ALIGN", (0, 1), (0, -1), "CENTER"),
                ("ALIGN", (4, 1), (4, -1), "CENTER"),
                ("ALIGN", (6, 1), (6, -1), "CENTER"),
                ("ALIGN", (8, 1), (8, -1), "CENTER"),
                ("VALIGN", (0, 1), (-1, -1), "TOP"),
            ]
            for i, _ in enumerate(chunk, start=1):
                src = comparison.rows[start + i - 1]
                bg = (
                    OK_BG
                    if src.verdict == "AGREE"
                    else FAIL_BG
                    if src.verdict == "DISAGREE"
                    else LIGHT
                )
                style_cmds.append(("BACKGROUND", (0, i), (-1, i), bg))
            table.setStyle(TableStyle(style_cmds))
            story.append(table)
            end = min(start + chunk_size, len(body_rows))
            story.append(
                Paragraph(f"Lines {start + 1}–{end} of {total_lines}", styles["tiny"])
            )
            story.append(Spacer(1, 3))
            if end < len(body_rows):
                story.append(PageBreak())
                story.append(
                    Paragraph(
                        f"Comparison table — continued ({end + 1}–{min(end + chunk_size, total_lines)} of {total_lines})",
                        styles["h"],
                    )
                )

    doc.build(story, onFirstPage=_header_footer, onLaterPages=_header_footer)
    return path
