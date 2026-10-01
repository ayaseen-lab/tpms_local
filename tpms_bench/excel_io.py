"""Copy the Hamaton database workbook and write bench result columns."""

from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

RESULT_COLUMNS = [
    "Board performance",
    "NOK reason",
    "Battery percentage",
    "Baterry voltage",
    "SDR compare",
    "SDR reason",
    "rtl_433 Decoder",
    "IQ file",
    "Sensor ID",
    "Frequency",
    "Pressure",
    "Temperature",
    "RSSI",
    "Duration s",
]


DATABASE_HEADERS = [
    "Make",
    "Model",
    "Year From",
    "Year To",
    "Region",
    "Platform",
    "Type",
    "OE supplier",
    "OE Number",
    "Freq",
    "Col11",
    "Col12",
    "Col13",
    "Col14",
    "Col15",
    "CODEA",
    "CODEB",
    "CODEC",
]


def _is_permission_error(exc: BaseException) -> bool:
    if isinstance(exc, PermissionError):
        return True
    if isinstance(exc, OSError) and getattr(exc, "errno", None) == 13:
        return True
    text = str(exc).lower()
    return "permission denied" in text or "being used by another process" in text


def is_valid_xlsx(path: Path) -> bool:
    """True when path is a readable OOXML workbook (ZIP-based .xlsx)."""
    try:
        if not path.is_file() or path.stat().st_size < 64:
            return False
        with path.open("rb") as fh:
            sig = fh.read(4)
        # ZIP local file header — all real .xlsx files start with PK\x03\x04
        if sig != b"PK\x03\x04":
            return False
        load_workbook(path)
        return True
    except Exception:
        return False


def alternate_workbook_path(preferred: Path) -> Path:
    """Unlocked sibling path when Excel (or another app) holds the preferred file."""
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return preferred.with_name(f"{preferred.stem}_live_{stamp}{preferred.suffix}")


def save_workbook(wb: Workbook, path: Path) -> Path:
    """Save workbook; if the target is locked (Excel open), write a live alternate."""
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(path)
        return path
    except Exception as exc:
        if not _is_permission_error(exc):
            raise
        alt = alternate_workbook_path(path)
        wb.save(alt)
        return alt


def copy_workbook(source: Path, dest: Path, *, force: bool = False) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    # Corrupt / truncated results files look like .xlsx but are not ZIP —
    # openpyxl then raises "File is not a zip file". Recreate from source.
    need_copy = force or not dest.exists() or not is_valid_xlsx(dest)
    if need_copy:
        try:
            shutil.copy2(source, dest)
            if not is_valid_xlsx(dest):
                raise RuntimeError(f"Copied workbook is unreadable: {dest}")
            return dest
        except Exception as exc:
            if not _is_permission_error(exc):
                if dest.exists() and not is_valid_xlsx(dest):
                    alt = alternate_workbook_path(dest)
                    shutil.copy2(source, alt)
                    return alt
                raise
            alt = alternate_workbook_path(dest)
            shutil.copy2(source, alt)
            return alt
    # Resume path: preferred exists but may be locked for later saves — probe write access.
    try:
        with dest.open("a+b"):
            pass
        return dest
    except Exception as exc:
        if not _is_permission_error(exc):
            raise
        alt = alternate_workbook_path(dest)
        shutil.copy2(source, alt)
        return alt


def vehicle_row_count(path: Path) -> int:
    """Number of catalog/data rows (header excluded)."""
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        n = max(0, int(ws.max_row or 1) - 1)
        wb.close()
        return n
    except Exception:
        return 0


def needs_catalog_refresh(source: Path, dest: Path) -> bool:
    """True when results workbook is empty but the selected catalog is full."""
    if not source.is_file() or not dest.is_file():
        return False
    src_n = vehicle_row_count(source)
    if src_n < 5:
        return False
    dest_n = vehicle_row_count(dest)
    return dest_n < 5


def create_blank_database(dest: Path) -> Path:
    """Minimal workbook so a run can start from custom opcodes only."""
    dest.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "TPMS Board"
    for col, name in enumerate(DATABASE_HEADERS, start=1):
        if not name.startswith("Col"):
            ws.cell(1, col, name)
    return save_workbook(wb, dest)


def append_manual_code_row(
    ws: Worksheet,
    cols: dict[str, int],
    *,
    code_a: str,
    code_b: str,
    code_c: str,
    label: str = "Manual code",
    make: str = "",
    model: str = "",
    freq: str = "",
) -> int:
    row = int(ws.max_row or 1) + 1
    if row < 2:
        row = 2
    # If the sheet only has a header, openpyxl may still report max_row=1.
    if ws.max_row == 1 and ws.cell(1, 1).value is None:
        row = 2
    make_text = (make or "").strip() or label or "Manual"
    model_text = (model or "").strip() or "Custom opcode"
    # Empty freq left SDR on whatever band was last used — default EU TPMS.
    freq_text = (freq or "").strip() or "433.92"
    ws.cell(row, cols.get("Make", 1), make_text)
    ws.cell(row, cols.get("Model", 2), model_text)
    ws.cell(row, cols.get("Year From", 3), "")
    ws.cell(row, cols.get("OE supplier", 8), "Manual")
    ws.cell(row, cols.get("OE Number", 9), label or "")
    ws.cell(row, cols.get("Freq", 10), freq_text)
    ws.cell(row, cols.get("CODEA", 16), code_a)
    ws.cell(row, cols.get("CODEB", 17), code_b)
    ws.cell(row, cols.get("CODEC", 18), code_c)
    return row


def manual_code_present(
    ws: Worksheet,
    cols: dict[str, int],
    code_a: str,
    code_b: str,
    code_c: str,
) -> bool:
    """True when a data row already has this CODE A/B/C triple."""
    ca = str(code_a or "").strip().upper()
    cb = str(code_b or "").strip().upper()
    cc = str(code_c or "").strip().upper()
    if not (ca and cb and cc):
        return False
    col_a = cols.get("CODEA", 16)
    col_b = cols.get("CODEB", 17)
    col_c = cols.get("CODEC", 18)
    for row in range(2, int(ws.max_row or 1) + 1):
        a = str(ws.cell(row, col_a).value or "").strip().upper()
        b = str(ws.cell(row, col_b).value or "").strip().upper()
        c = str(ws.cell(row, col_c).value or "").strip().upper()
        if a == ca and b == cb and c == cc:
            return True
    return False


def _header_map(ws: Worksheet) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for cell in ws[1]:
        if cell.value:
            mapping[str(cell.value).strip()] = cell.column
    return mapping


def ensure_result_columns(ws: Worksheet) -> dict[str, int]:
    mapping = _header_map(ws)
    last = max(mapping.values()) if mapping else 1
    for name in RESULT_COLUMNS:
        if name not in mapping:
            last += 1
            ws.cell(1, last, name)
            mapping[name] = last
    return mapping


def parse_code(value) -> int | None:
    if value is None or value == "":
        return None
    text = str(value).strip()
    if text.lower() in ("none", "nan"):
        return None
    try:
        return int(text, 16)
    except ValueError:
        try:
            return int(float(text))
        except ValueError:
            return None


def load_output(path: Path) -> tuple[Workbook, Worksheet, dict[str, int]]:
    if not is_valid_xlsx(path):
        raise RuntimeError(
            f"Results Excel is damaged or not a real .xlsx file:\n{path}\n\n"
            "Close Excel if it is open, then Start again — a fresh copy will be created."
        )
    wb = load_workbook(path)
    ws = wb[wb.sheetnames[0]]
    cols = ensure_result_columns(ws)
    return wb, ws, cols


def stamp_board_report_info(path: Path) -> None:
    """Mark the workbook as a TPMS Board report so it is not confused with SDR exports."""
    if not is_valid_xlsx(path):
        return
    wb = load_workbook(path)
    if "Report Info" in wb.sheetnames:
        ws = wb["Report Info"]
    else:
        ws = wb.create_sheet("Report Info", 1)
    ws["A1"] = "TPMS Board Report"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Hamaton board validation results — this is not an SDR receiver report"
    ws["A4"] = "Report type"
    ws["B4"] = "TPMS Board (USB-TTL / Hamaton bench) — not SDR Receiver"
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 64
    save_workbook(wb, path)


def write_row(ws: Worksheet, cols: dict[str, int], row: int, values: dict[str, object]) -> None:
    # Auto-append any new result keys (e.g. RSSI) so write never KeyErrors.
    last = max(cols.values()) if cols else 1
    for name in values:
        if name not in cols:
            last += 1
            ws.cell(1, last, name)
            cols[name] = last
    for name, value in values.items():
        ws.cell(row, cols[name], value)
