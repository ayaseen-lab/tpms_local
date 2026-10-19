"""TPMS Board view — Hamaton bench UI as a CustomTkinter frame."""

from __future__ import annotations

import os
import platform
import queue
import shutil
import subprocess
import threading
import time
import tkinter as tk
from datetime import datetime
from pathlib import Path
from tkinter import messagebox, ttk
from typing import Callable, Optional

import customtkinter as ctk

from dialogs import ask_open_path, ask_save_path
from themes import (
    COLOR_BG,
    COLOR_BG_CARD,
    COLOR_BG_PANEL,
    COLOR_BLUE,
    COLOR_BLUE_BG,
    COLOR_BORDER,
    COLOR_BTN_EXPORT,
    COLOR_BTN_EXPORT_HOVER,
    COLOR_BTN_PAUSE,
    COLOR_BTN_PAUSE_HOVER,
    COLOR_BTN_PRIMARY,
    COLOR_BTN_PRIMARY_HOVER,
    COLOR_BTN_PRIMARY_TEXT,
    COLOR_BTN_SECONDARY,
    COLOR_BTN_SECONDARY_HOVER,
    COLOR_BTN_STOP_BORDER,
    COLOR_BTN_STOP_OUTLINE,
    COLOR_BTN_STOP_OUTLINE_HOVER,
    COLOR_BTN_STOP_TEXT,
    COLOR_CYAN_BG,
    COLOR_GREEN,
    COLOR_GREEN_BG,
    COLOR_HEADER_ACCENT,
    COLOR_HEADER_BG,
    COLOR_LIVE_BG,
    COLOR_LIVE_CHIP,
    COLOR_LIVE_FG,
    COLOR_LIVE_MUTED,
    COLOR_LIVE_TEAL,
    COLOR_ORANGE,
    COLOR_ORANGE_BG,
    COLOR_PILL_OK_BG,
    COLOR_PILL_OK_FG,
    COLOR_RED,
    COLOR_RED_BG,
    COLOR_TEXT,
    COLOR_TEXT_DIM,
    COLOR_TEXT_MUTED,
    combo_colors,
    entry_colors,
    ui_font,
)
from tpms_bench.excel_io import parse_code, stamp_board_report_info, vehicle_row_count
from tpms_bench.paths import app_root
from tpms_bench.report import build_pdf
from tpms_bench.runner import OUT_XLSX, SPEC_XLSX, BenchRunner, ManualCode, ProgressEvent, get_out_xlsx, reset_session_db
from tpms_bench.uart import default_port, find_serial_ports, port_device
from activity_terminal import ActivityTerminalPanel
from charts import BoardTrendCharts, ResultCharts
from widgets import StatCard

CODE_A_BG = "#E7F4F7"
CODE_A_FG = "#1B5F70"
CODE_B_BG = "#F8EFE6"
CODE_B_FG = "#8A4A1C"
CODE_C_BG = "#E8F7F4"
CODE_C_FG = "#14685E"
TELEMETRY_BG = "#E7F6EF"
TELEMETRY_FG = "#1A8F6E"
LED_IDLE = "#64748B"
LED_TTL_ON = "#34D399"
LED_JLINK_ON = "#FBBF24"

# In-repo Hamaton catalog — always loaded from data/ (tracked in git).
DEFAULT_HAMATON_XLSX = SPEC_XLSX if SPEC_XLSX.is_file() else (app_root() / "data" / "Hamaton_database_20260126_1305.xlsx")


def _fmt_batt(voltage: str | None, percentage: str | None = None) -> str:
    """Pretty battery for live labels — never show '0%' or bare 'na V'."""
    v = str(voltage or "").strip()
    p = str(percentage or "").strip()
    if v and v.lower() not in {"na", "n/a", "—", "-", "0", "0.0", "0.000"}:
        if v.upper() in {"OK", "LOW"}:
            return f"Batt {v}"
        if v.endswith("%"):
            return f"Batt {v}"
        try:
            return f"{float(v):.2f} V"
        except ValueError:
            return f"Batt {v}"
    if p and p.lower() not in {"na", "n/a", "—", "-", "0"}:
        try:
            n = int(float(p))
            if 1 <= n <= 100:
                return f"Batt {n}%"
        except ValueError:
            if p.endswith("%"):
                return f"Batt {p}"
    return "Batt —"


class _ControlState:
    """Stand-in for hidden RUN CHUNK / RUN FULL so Pause/Stop still track session state."""

    def __init__(self) -> None:
        self._state = "normal"

    def configure(self, **kwargs) -> None:
        if "state" in kwargs:
            self._state = kwargs["state"]

    def cget(self, key: str):
        if key == "state":
            return self._state
        return ""


class TpmsView(ctk.CTkFrame):
    """Board validation UI. Switching away from this frame does not stop the bench."""

    def __init__(
        self,
        master,
        on_status_change: Optional[Callable[[bool, str], None]] = None,
        on_chunk_complete: Optional[Callable[[str], None]] = None,
        on_process_update: Optional[Callable[..., None]] = None,
        chunk_var: Optional[tk.StringVar] = None,
        **kwargs,
    ):
        kwargs.setdefault("fg_color", COLOR_BG)
        super().__init__(master, **kwargs)
        self.on_status_change = on_status_change or (lambda _running, _label: None)
        self.on_chunk_complete = on_chunk_complete or (lambda _msg: None)
        self.on_process_update = on_process_update or (lambda *_a, **_k: None)
        self.on_board_rf: Callable | None = None
        self.on_board_trigger: Callable | None = None
        self.on_sdr_finalize: Callable | None = None
        # Combined shell: ensure SDR is listening (and tuned) before/during Board rows.
        self.on_prepare_sdr: Callable[..., bool] | None = None
        # Combined shell: release/stop SDR after Board session/chunk settles.
        self.on_sdr_release: Callable[..., None] | None = None
        # Optional sync prepare used by BenchRunner (blocks until tuned).
        self.live_sdr_prepare: Callable[..., bool] | None = None
        self.chunk_var = chunk_var or tk.StringVar(value="100")

        self.queue: queue.Queue[ProgressEvent] = queue.Queue()
        self.runner: BenchRunner | None = None
        self.worker: threading.Thread | None = None
        self.is_paused = False
        self.start_time: float | None = None
        self.source_xlsx: Path | None = None
        self._total_rows = 0
        self._anim_step = 0
        self._state_label = "IDLE"
        self.custom_codes: list[ManualCode] = []
        self._sdr_agree = 0
        self._sdr_disagree = 0
        # Current Board session only — used by Comparative export (not stale SQLite history).
        self._session_rows: list[dict] = []
        self._last_board_tel: dict | None = None

        self._setup_styles()
        self._build_ui()
        self._reset_display()
        # Session restore is deferred — CombinedApp asks Continue vs New Session at startup.
        self._autoload_default_excel()
        self.after(150, self._drain)
        self.after(80, self._animate_badge)
        self.after(900, self._poll_live_sdr)
        self._notify_status()
        self.live_sdr_get = None
        # Board-only sessions (USB-TTL) never prepare / wait on SDR.
        self._session_skip_sdr = True

    def _toplevel(self):
        return self.winfo_toplevel()

    def _sdr_enabled(self) -> bool:
        """False for Board-only USB-TTL runs — TPMS Board and SDR stay separate."""
        return not bool(getattr(self, "_session_skip_sdr", True))

    def _setup_styles(self) -> None:
        style = ttk.Style(self._toplevel())
        style.theme_use("clam")
        style.configure(
            "Combined.Treeview",
            background=COLOR_BG_CARD,
            foreground=COLOR_TEXT,
            fieldbackground=COLOR_BG_CARD,
            rowheight=24,
            font=("Segoe UI", 10),
            bordercolor=COLOR_BORDER,
            borderwidth=0,
        )
        style.configure(
            "Combined.Treeview.Heading",
            background="#0B0F14",
            foreground="#ffffff",
            font=("Segoe UI", 9, "bold"),
            relief="flat",
            padding=6,
        )
        style.map(
            "Combined.Treeview.Heading",
            background=[("active", "#12263A")],
            foreground=[("active", "#ffffff")],
        )
        style.map(
            "Combined.Treeview",
            background=[("selected", "#D6E8F5")],
            foreground=[("selected", COLOR_TEXT)],
            fieldbackground=[("", COLOR_BG)],
        )
        style.configure(
            "Combined.Vertical.TScrollbar",
            background="#C5CED8",
            troughcolor=COLOR_BG,
            arrowcolor=COLOR_HEADER_BG,
            bordercolor=COLOR_BORDER,
        )
        style.configure(
            "Combined.Horizontal.TScrollbar",
            background="#C5CED8",
            troughcolor=COLOR_BG,
            arrowcolor=COLOR_HEADER_BG,
            bordercolor=COLOR_BORDER,
        )

    def _build_ui(self) -> None:
        """Scrollable Board page: live telemetry + session + full history."""
        LIVE_BG = COLOR_LIVE_BG
        LIVE_FG = COLOR_LIVE_FG
        LIVE_MUTED = COLOR_LIVE_MUTED
        LIVE_CHIP = COLOR_LIVE_CHIP
        LIVE_TEAL = COLOR_LIVE_TEAL

        self.grid_columnconfigure(0, weight=1)
        self.grid_rowconfigure(0, weight=1)

        # Whole Board view scrolls so history is always reachable.
        self._page = ctk.CTkScrollableFrame(
            self, fg_color="transparent", corner_radius=0,
            scrollbar_button_color=COLOR_BORDER,
            scrollbar_button_hover_color=COLOR_TEXT_MUTED,
        )
        self._page.grid(row=0, column=0, sticky="nsew")
        self._page.grid_columnconfigure(0, weight=1)
        root = self._page

        # —— Metric cards ——
        stats_row = ctk.CTkFrame(root, fg_color="transparent")
        stats_row.grid(row=0, column=0, sticky="ew", pady=(0, 10))
        for i in range(5):
            stats_row.grid_columnconfigure(i, weight=1, minsize=120)
        self.stat_pending = StatCard(
            stats_row, "Pending", "—", COLOR_ORANGE, COLOR_ORANGE_BG, compact=True, subtitle="in queue"
        )
        self.stat_pending.grid(row=0, column=0, sticky="ew", padx=(0, 8))
        self.stat_done = StatCard(
            stats_row, "Tested", "0", COLOR_BLUE, COLOR_BLUE_BG, compact=True, subtitle="of 0 rows"
        )
        self.stat_done.grid(row=0, column=1, sticky="ew", padx=(0, 8))
        self.stat_ok = StatCard(
            stats_row, "Passed (OK)", "0", COLOR_GREEN, COLOR_GREEN_BG, compact=True, subtitle="all codes matched"
        )
        self.stat_ok.grid(row=0, column=2, sticky="ew", padx=(0, 8))
        self.stat_nok = StatCard(
            stats_row, "Failed (NOK)", "0", COLOR_RED, COLOR_RED_BG, compact=True, subtitle="no failures"
        )
        self.stat_nok.grid(row=0, column=3, sticky="ew", padx=(0, 8))
        self.stat_rate = StatCard(
            stats_row, "Pass rate", "0.0%", COLOR_GREEN, COLOR_GREEN_BG, compact=True, subtitle=""
        )
        self.stat_rate.grid(row=0, column=4, sticky="ew")

        # —— Mid: live (left, fills height) + session (right) ——
        # Left column stretches to match the taller session sidebar so
        # telemetry occupies the space that used to look empty.
        mid = ctk.CTkFrame(root, fg_color="transparent")
        mid.grid(row=1, column=0, sticky="ew", pady=(0, 10))
        mid.grid_columnconfigure(0, weight=3)
        mid.grid_columnconfigure(1, weight=2)
        mid.grid_rowconfigure(0, weight=1)

        # Live dark card — fills mid height; telemetry expands into leftover space
        live = ctk.CTkFrame(mid, fg_color=LIVE_BG, corner_radius=16)
        live.grid(row=0, column=0, sticky="nsew", padx=(0, 10))
        live_in = ctk.CTkFrame(live, fg_color="transparent")
        live_in.pack(fill="both", expand=True, padx=14, pady=12)

        head = ctk.CTkFrame(live_in, fg_color="transparent")
        head.pack(fill="x")
        self.state_badge = ctk.CTkLabel(
            head,
            text="IDLE",
            font=ui_font(11, "bold"),
            text_color=COLOR_PILL_OK_FG,
            fg_color=COLOR_PILL_OK_BG,
            corner_radius=8,
            height=26,
            width=90,
        )
        self.state_badge.pack(side="left", padx=(0, 10))
        self.row_counter = ctk.CTkLabel(
            head, text="ROW —", font=ui_font(12, "bold"), text_color=LIVE_MUTED
        )
        self.row_counter.pack(side="right")

        self.vehicle_var = ctk.StringVar(value="Ready — choose Excel or custom codes")
        ctk.CTkLabel(
            live_in,
            textvariable=self.vehicle_var,
            font=ui_font(20, "bold"),
            text_color=LIVE_FG,
            anchor="w",
        ).pack(fill="x", pady=(8, 6))

        self.meta_var = ctk.StringVar(value="OE: —   ·   Supplier: —   ·   Freq: —")
        meta_row = ctk.CTkFrame(live_in, fg_color="transparent")
        meta_row.pack(fill="x", pady=(0, 10))
        self.meta_chip_oe = ctk.CTkLabel(
            meta_row, text="OE —", font=ui_font(11, "bold"), text_color=LIVE_MUTED,
            fg_color=LIVE_CHIP, corner_radius=8, height=26, padx=10,
        )
        self.meta_chip_oe.pack(side="left", padx=(0, 6))
        self.meta_chip_supplier = ctk.CTkLabel(
            meta_row, text="Supplier —", font=ui_font(11, "bold"), text_color=LIVE_MUTED,
            fg_color=LIVE_CHIP, corner_radius=8, height=26, padx=10,
        )
        self.meta_chip_supplier.pack(side="left", padx=(0, 6))
        self.meta_chip_freq = ctk.CTkLabel(
            meta_row, text="Freq —", font=ui_font(11, "bold"), text_color=LIVE_MUTED,
            fg_color=LIVE_CHIP, corner_radius=8, height=26, padx=10,
        )
        self.meta_chip_freq.pack(side="left")
        # Keep meta_var for existing setters; sync chips via trace.
        self.meta_var.trace_add("write", lambda *_: self._sync_meta_chips())

        codes = ctk.CTkFrame(live_in, fg_color="transparent")
        codes.pack(fill="x", pady=(0, 10))
        self.code_vars = {"A": ctk.StringVar(value="—"), "B": ctk.StringVar(value="—"), "C": ctk.StringVar(value="—")}
        for key, bg, fg, pad in (
            ("A", "#1E3A42", "#7DD3E8", 8),
            ("B", "#3A2E1E", "#E8C07D", 8),
            ("C", "#1E3A32", "#7DE8C0", 0),
        ):
            box = ctk.CTkFrame(codes, fg_color=bg, corner_radius=12)
            box.pack(side="left", expand=True, fill="both", padx=(0, pad))
            inner = ctk.CTkFrame(box, fg_color="transparent")
            inner.pack(fill="x", padx=12, pady=10)
            topc = ctk.CTkFrame(inner, fg_color="transparent")
            topc.pack(fill="x")
            ctk.CTkLabel(topc, text=f"CODE {key}", font=ui_font(10, "bold"), text_color=fg).pack(side="left")
            ctk.CTkButton(
                topc, text="⎘", width=28, height=22, font=ui_font(12),
                fg_color="transparent", hover_color=LIVE_CHIP, text_color=fg,
                command=lambda k=key: self._copy_code(k),
            ).pack(side="right")
            ctk.CTkLabel(
                inner, textvariable=self.code_vars[key],
                font=ctk.CTkFont(family="Consolas", size=16, weight="bold"),
                text_color=fg, anchor="w",
            ).pack(anchor="w", pady=(4, 0))

        # —— Board live telemetry fills remaining height under codes ——
        tel = ctk.CTkFrame(live_in, fg_color=LIVE_CHIP, corner_radius=12)
        tel.pack(fill="both", expand=True, pady=(0, 0))
        tel_in = ctk.CTkFrame(tel, fg_color="transparent")
        tel_in.pack(fill="both", expand=True, padx=12, pady=10)
        tel_head = ctk.CTkFrame(tel_in, fg_color="transparent")
        tel_head.pack(fill="x", pady=(0, 8))
        ctk.CTkLabel(
            tel_head,
            text="Board live telemetry",
            font=ui_font(12, "bold"),
            text_color=LIVE_TEAL,
        ).pack(side="left")
        self.tel_phase_var = ctk.StringVar(value="Waiting for Board read…")
        ctk.CTkLabel(
            tel_head,
            textvariable=self.tel_phase_var,
            font=ui_font(11),
            text_color=LIVE_MUTED,
        ).pack(side="right")

        # Horizontal scroll so five metric cells stay readable in a narrow column
        tel_scroll = ctk.CTkScrollableFrame(
            tel_in,
            fg_color="transparent",
            corner_radius=0,
            orientation="horizontal",
            height=150,
            scrollbar_button_color="#2A3340",
            scrollbar_button_hover_color=LIVE_MUTED,
        )
        tel_scroll.pack(fill="both", expand=True, pady=(0, 4))

        self.tel_id = ctk.StringVar(value="—")
        self.tel_temp = ctk.StringVar(value="—")
        self.tel_pressure = ctk.StringVar(value="—")
        self.tel_rssi = ctk.StringVar(value="—")
        self.tel_batt = ctk.StringVar(value="—")
        self.reading_var = ctk.StringVar(value="Waiting for first reading…")
        self._tel_value_labels: dict[str, ctk.CTkLabel] = {}

        def _tel_cell(key: str, title: str, var: ctk.StringVar, accent: str = LIVE_FG, last: bool = False):
            cell = ctk.CTkFrame(
                tel_scroll,
                fg_color=LIVE_BG,
                corner_radius=10,
                border_width=1,
                border_color="#2A3340",
                width=148,
            )
            cell.pack(side="left", fill="y", expand=True, padx=(0 if key == "id" else 6, 0 if last else 0))
            cell.pack_propagate(False)
            inner = ctk.CTkFrame(cell, fg_color="transparent")
            inner.pack(fill="both", expand=True, padx=12, pady=12)
            ctk.CTkLabel(inner, text=title, font=ui_font(10, "bold"), text_color=LIVE_MUTED).pack(
                anchor="w"
            )
            val = ctk.CTkLabel(
                inner,
                textvariable=var,
                font=ctk.CTkFont(family="Consolas", size=18, weight="bold"),
                text_color=accent,
                anchor="w",
                wraplength=130,
                justify="left",
            )
            val.pack(anchor="w", pady=(8, 0), fill="x")
            self._tel_value_labels[key] = val

        _tel_cell("id", "Sensor ID", self.tel_id, LIVE_TEAL)
        _tel_cell("temp", "Temperature", self.tel_temp)
        _tel_cell("pressure", "Pressure", self.tel_pressure)
        _tel_cell("rssi", "RSSI", self.tel_rssi)
        _tel_cell("batt", "Battery", self.tel_batt, last=True)

        # Summary line uses leftover vertical space under the metric strip
        self.tel_summary = ctk.CTkLabel(
            tel_in,
            textvariable=self.reading_var,
            font=ui_font(12),
            text_color=LIVE_MUTED,
            anchor="w",
            justify="left",
        )
        self.tel_summary.pack(fill="x", pady=(10, 0))

        # Right: session setup + custom codes (content-sized, no stretch)
        side = ctk.CTkFrame(mid, fg_color="transparent")
        side.grid(row=0, column=1, sticky="nsew")

        setup = ctk.CTkFrame(
            side, fg_color=COLOR_BG_CARD, corner_radius=14, border_width=1, border_color=COLOR_BORDER
        )
        setup.pack(fill="x", pady=(0, 10))
        setup_in = ctk.CTkFrame(setup, fg_color="transparent")
        setup_in.pack(fill="x", padx=14, pady=12)

        setup_head = ctk.CTkFrame(setup_in, fg_color="transparent")
        setup_head.pack(fill="x", pady=(0, 10))
        ctk.CTkLabel(
            setup_head, text="Session setup", font=ui_font(14, "bold"), text_color=COLOR_TEXT
        ).pack(side="left")
        ctk.CTkButton(
            setup_head, text="Reset session", width=110, height=28, font=ui_font(11),
            fg_color=COLOR_BG_PANEL, hover_color=COLOR_BORDER, text_color=COLOR_TEXT, corner_radius=8,
            command=self.reset_session,
        ).pack(side="right")

        ctk.CTkLabel(
            setup_in, text="USB serial port", font=ui_font(10, "bold"), text_color=COLOR_TEXT_MUTED
        ).pack(anchor="w")
        port_row = ctk.CTkFrame(setup_in, fg_color="transparent")
        port_row.pack(fill="x", pady=(4, 10))
        self.port_var = ctk.StringVar(value=default_port())
        self.port_combo = ctk.CTkComboBox(
            port_row, values=find_serial_ports() or [default_port()],
            width=220, height=32, font=ui_font(12), **combo_colors(),
        )
        self.port_combo.set(self.port_var.get())
        self.port_combo.pack(side="left", fill="x", expand=True)
        ctk.CTkButton(
            port_row, text="↻", width=32, height=32, font=ui_font(13, "bold"),
            fg_color=COLOR_BG_PANEL, hover_color=COLOR_BORDER, text_color=COLOR_TEXT, corner_radius=8,
            command=self._refresh_ports,
        ).pack(side="left", padx=(6, 0))

        ctk.CTkLabel(
            setup_in, text="Firmware", font=ui_font(10, "bold"), text_color=COLOR_TEXT_MUTED
        ).pack(anchor="w")
        fw_row = ctk.CTkFrame(setup_in, fg_color="transparent")
        fw_row.pack(fill="x", pady=(4, 10))
        self.version_var = ctk.StringVar(value="—")
        ctk.CTkLabel(
            fw_row, textvariable=self.version_var, font=ui_font(12, "bold"), text_color=COLOR_TEXT, anchor="w"
        ).pack(side="left", fill="x", expand=True)
        self.query_fw_btn = ctk.CTkButton(
            fw_row, text="Details", width=72, height=28, font=ui_font(11, "bold"),
            fg_color=COLOR_BG_CARD, hover_color=COLOR_BG_PANEL, text_color=COLOR_TEXT,
            border_width=1, border_color=COLOR_BORDER, corner_radius=8,
            command=self.query_firmware_version,
        )
        self.query_fw_btn.pack(side="right")

        ctk.CTkLabel(
            setup_in, text="Excel database", font=ui_font(10, "bold"), text_color=COLOR_TEXT_MUTED
        ).pack(anchor="w")
        self.file_label = ctk.CTkLabel(
            setup_in, text="No file selected", font=ui_font(12, "bold"), text_color=COLOR_ORANGE, anchor="w"
        )
        self.file_label.pack(fill="x", pady=(4, 6))
        excel_btns = ctk.CTkFrame(setup_in, fg_color="transparent")
        excel_btns.pack(fill="x")
        ctk.CTkButton(
            excel_btns, text="Change file", width=100, height=30, font=ui_font(11, "bold"),
            fg_color=COLOR_BG_CARD, hover_color=COLOR_BG_PANEL, text_color=COLOR_TEXT,
            border_width=1, border_color=COLOR_BORDER, corner_radius=8, command=self.select_excel,
        ).pack(side="left")
        self.remove_excel_btn = ctk.CTkButton(
            excel_btns, text="Remove", width=80, height=30, font=ui_font(11),
            fg_color=COLOR_BG_CARD, hover_color=COLOR_RED_BG, text_color=COLOR_BTN_STOP_TEXT,
            border_width=1, border_color=COLOR_BTN_STOP_BORDER, corner_radius=8,
            command=self.remove_excel, state="disabled",
        )
        self.remove_excel_btn.pack(side="left", padx=(8, 0))

        # Hidden control refs used by start/pause/stop state machine
        ctrl = ctk.CTkFrame(setup_in, fg_color="transparent")
        ctrl.pack(fill="x", pady=(12, 0))
        self.full_btn = _ControlState()
        self.transport_var = ctk.StringVar(value="USB-TTL TX/RX")
        self.start_btn = ctk.CTkButton(
            ctrl, text="Start Board", width=110, height=34, font=ui_font(12, "bold"),
            fg_color=COLOR_BTN_PRIMARY, hover_color=COLOR_BTN_PRIMARY_HOVER,
            text_color=COLOR_BTN_PRIMARY_TEXT, corner_radius=10,
            command=lambda: self.start_test(skip_sdr=True, run_full=False),
        )
        self.start_btn.pack(side="left", padx=(0, 6))
        self.pause_btn = ctk.CTkButton(
            ctrl, text="Pause test", width=100, height=34, font=ui_font(12, "bold"),
            fg_color=COLOR_BTN_PAUSE, hover_color=COLOR_BTN_PAUSE_HOVER, text_color="#FFFFFF",
            corner_radius=10, state="disabled", command=self.toggle_pause,
        )
        self.pause_btn.pack(side="left", padx=(0, 6))
        self.stop_btn = ctk.CTkButton(
            ctrl, text="Stop", width=72, height=34, font=ui_font(12, "bold"),
            fg_color=COLOR_BTN_STOP_OUTLINE, hover_color=COLOR_BTN_STOP_OUTLINE_HOVER,
            text_color=COLOR_BTN_STOP_TEXT, border_width=2, border_color=COLOR_BTN_STOP_BORDER,
            corner_radius=10, state="disabled", command=self.stop_test,
        )
        self.stop_btn.pack(side="left")

        custom = ctk.CTkFrame(
            side, fg_color=COLOR_BG_CARD, corner_radius=14, border_width=1, border_color=COLOR_BORDER
        )
        custom.pack(fill="x", expand=False)
        custom_in = ctk.CTkFrame(custom, fg_color="transparent")
        custom_in.pack(fill="x", expand=False, padx=12, pady=10)
        ctk.CTkLabel(
            custom_in, text="Custom codes", font=ui_font(13, "bold"), text_color=COLOR_TEXT
        ).pack(anchor="w", pady=(0, 6))
        self.manual_entries: dict[str, ctk.CTkEntry] = {}
        for key in ("A", "B", "C"):
            row = ctk.CTkFrame(custom_in, fg_color="transparent")
            row.pack(fill="x", pady=2)
            ctk.CTkLabel(row, text=f"Code {key}", width=58, anchor="w", font=ui_font(11), text_color=COLOR_TEXT_MUTED).pack(side="left")
            entry = ctk.CTkEntry(row, height=30, placeholder_text=f"{key} hex", **entry_colors())
            entry.pack(side="left", fill="x", expand=True)
            entry.bind("<Return>", lambda _e: self.add_custom_code())
            self.manual_entries[key] = entry
        meta = ctk.CTkFrame(custom_in, fg_color="transparent")
        meta.pack(fill="x", pady=(8, 0))
        self.manual_make_entry = ctk.CTkEntry(meta, width=90, height=28, placeholder_text="Make", **entry_colors())
        self.manual_make_entry.pack(side="left", padx=(0, 4))
        self.manual_model_entry = ctk.CTkEntry(meta, width=90, height=28, placeholder_text="Model", **entry_colors())
        self.manual_model_entry.pack(side="left", padx=(0, 4))
        self.manual_freq_entry = ctk.CTkEntry(meta, width=70, height=28, placeholder_text="MHz", **entry_colors())
        self.manual_freq_entry.pack(side="left", padx=(0, 4))
        self.manual_label_entry = ctk.CTkEntry(meta, width=80, height=28, placeholder_text="Label", **entry_colors())
        self.manual_label_entry.pack(side="left", padx=(0, 4))
        self.manual_label_entry.bind("<Return>", lambda _e: self.add_custom_code())
        ctk.CTkButton(
            custom_in, text="Add code", width=100, height=30, font=ui_font(11, "bold"),
            fg_color=COLOR_GREEN, hover_color="#15803d", corner_radius=8, command=self.add_custom_code,
        ).pack(anchor="w", pady=(8, 0))
        self.codes_list = ctk.CTkFrame(custom_in, fg_color="transparent")
        self.codes_list.pack(fill="x", pady=(6, 0))
        self.codes_hint = ctk.CTkLabel(
            self.codes_list, text="No custom codes added yet.", font=ui_font(10),
            text_color=COLOR_TEXT_MUTED, anchor="w",
        )
        self.codes_hint.pack(anchor="w")

        # Charts kept for logic (hidden)
        charts_anchor = ctk.CTkFrame(root, fg_color="transparent", height=1)
        self.trend_charts = BoardTrendCharts(charts_anchor, height=1, title="")
        self.charts = ResultCharts(charts_anchor, height=1, title="")

        # —— Full-width progress (not beside Custom codes) ——
        prog_card = ctk.CTkFrame(
            root, fg_color=COLOR_LIVE_BG, corner_radius=12
        )
        prog_card.grid(row=2, column=0, sticky="ew", pady=(0, 10))
        prog_wrap = ctk.CTkFrame(prog_card, fg_color="transparent")
        prog_wrap.pack(fill="x", padx=14, pady=10)
        self.progress = ctk.CTkProgressBar(
            prog_wrap, height=8, progress_color=COLOR_LIVE_TEAL, fg_color="#2A313A", corner_radius=4
        )
        self.progress.pack(fill="x")
        self.progress.set(0)
        prog_row = ctk.CTkFrame(prog_wrap, fg_color="transparent")
        prog_row.pack(fill="x", pady=(6, 0))
        self.progress_label = ctk.CTkLabel(
            prog_row, text="Progress: 0.0% (0 of 0)", font=ui_font(11), text_color=COLOR_LIVE_MUTED
        )
        self.progress_label.pack(side="left")
        self.eta_label = ctk.CTkLabel(prog_row, text="", font=ui_font(11), text_color=COLOR_LIVE_MUTED)
        self.eta_label.pack(side="right")
        self.elapsed_label = ctk.CTkLabel(prog_row, text="", font=ui_font(11), text_color=COLOR_LIVE_MUTED)
        self.elapsed_label.pack(side="right", padx=(0, 12))
        self.speed_label = ctk.CTkLabel(prog_row, text="", font=ui_font(11), text_color=COLOR_LIVE_MUTED)
        self.speed_label.pack(side="right", padx=(0, 12))

        # —— TPMS Board matches (same visual style as SDR matches) ——
        results = tk.Frame(
            root, bg=COLOR_BG, highlightbackground=COLOR_BORDER, highlightthickness=1
        )
        results.grid(row=3, column=0, sticky="ew", pady=(0, 6))

        table_header = tk.Frame(results, bg=COLOR_BG_PANEL)
        table_header.pack(fill="x")
        tk.Label(
            table_header,
            text="TPMS Board matches — one row per Board test (USB-TTL telemetry)",
            bg=COLOR_BG_PANEL,
            fg=COLOR_TEXT,
            font=("Segoe UI", 10, "bold"),
        ).pack(side="left", padx=8, pady=4)
        self.table_count_label = tk.Label(
            table_header, text="0 tests", bg=COLOR_BG_PANEL, fg=COLOR_TEXT_DIM, font=("Segoe UI", 9)
        )
        self.table_count_label.pack(side="right", padx=8)

        tools = ctk.CTkFrame(results, fg_color=COLOR_BG_PANEL)
        tools.pack(fill="x", padx=4, pady=(0, 2))
        self.history_filter = ctk.CTkComboBox(
            tools,
            values=["All results", "Passed only", "Failed only"],
            width=130,
            height=26,
            font=ui_font(11),
            command=self._apply_history_filter,
            **combo_colors(),
        )
        self.history_filter.set("All results")
        self.history_filter.pack(side="right", padx=(0, 4), pady=4)
        self.history_search = ctk.CTkEntry(
            tools, width=160, height=26, placeholder_text="Search vehicle…", **entry_colors()
        )
        self.history_search.pack(side="right", padx=(0, 6), pady=4)
        self.history_search.bind("<KeyRelease>", lambda _e: self._apply_history_filter())
        ctk.CTkButton(
            tools, text="Export Excel", width=96, height=26, font=ui_font(11, "bold"),
            fg_color=COLOR_BTN_EXPORT, hover_color=COLOR_BTN_EXPORT_HOVER, corner_radius=6,
            command=self.export_excel,
        ).pack(side="right", padx=(0, 6), pady=4)
        ctk.CTkButton(
            tools, text="Export PDF", width=88, height=26, font=ui_font(11, "bold"),
            fg_color=COLOR_BTN_EXPORT, hover_color=COLOR_BTN_EXPORT_HOVER, corner_radius=6,
            command=self.export_pdf,
        ).pack(side="right", padx=(0, 4), pady=4)

        tree_host = tk.Frame(results, bg=COLOR_BG)
        tree_host.pack(fill="x", padx=4, pady=(0, 4))
        columns = (
            "num",
            "vehicle",
            "codes",
            "id",
            "result",
            "pressure",
            "pressure_bar",
            "temp",
            "battery",
            "rssi",
            "reads",
            "time",
            "reason",
        )
        self.tree = ttk.Treeview(
            tree_host, columns=columns, show="headings", height=10, style="Combined.Treeview"
        )
        headings = {
            "num": "#",
            "vehicle": "Vehicle",
            "codes": "CODE A / B / C",
            "id": "Sensor ID",
            "result": "Result",
            "pressure": "PSI",
            "pressure_bar": "bar",
            "temp": "°C",
            "battery": "Batt",
            "rssi": "RSSI",
            "reads": "Reads",
            "time": "Seen",
            "reason": "Notes / NOK reason",
        }
        self._tree_col_weights = {
            "num": 0.04,
            "vehicle": 0.12,
            "codes": 0.12,
            "id": 0.09,
            "result": 0.06,
            "pressure": 0.06,
            "pressure_bar": 0.05,
            "temp": 0.05,
            "battery": 0.06,
            "rssi": 0.06,
            "reads": 0.05,
            "time": 0.07,
            "reason": 0.17,
        }
        for col in columns:
            self.tree.heading(col, text=headings[col])
            self.tree.column(
                col,
                width=80,
                minwidth=36,
                stretch=True,
                anchor="center" if col not in ("reason", "vehicle", "codes") else "w",
            )
        scroll_y = ttk.Scrollbar(
            tree_host, orient="vertical", command=self.tree.yview, style="Combined.Vertical.TScrollbar"
        )
        self.tree.configure(yscrollcommand=scroll_y.set)
        tree_host.grid_columnconfigure(0, weight=1)
        self.tree.grid(row=0, column=0, sticky="ew")
        scroll_y.grid(row=0, column=1, sticky="ns")
        self.tree.tag_configure("OK", background=COLOR_GREEN_BG, foreground="#047857")
        self.tree.tag_configure("NOK", background=COLOR_RED_BG, foreground="#B91C1C")
        self.tree.tag_configure("SKIP", background=COLOR_BG_PANEL, foreground=COLOR_TEXT_MUTED)
        self.tree.tag_configure("flash", background=COLOR_CYAN_BG)
        self.tree.bind("<Configure>", self._fit_tree_columns, add="+")
        self.after(50, self._fit_tree_columns)
        self._history_all: list[tuple] = []
        self._board_read_counts: dict[str, int] = {}

        # Activity terminal (same style as SDR tab)
        self.activity_terminal = ActivityTerminalPanel(root, start_expanded=False)
        self.activity_terminal.grid(row=4, column=0, sticky="ew", pady=(4, 0))

        footer = ctk.CTkFrame(root, fg_color="transparent")
        footer.grid(row=5, column=0, sticky="ew", pady=(4, 8))
        self.status_var = ctk.StringVar(
            value="Ready · USB-TTL Board (SDR optional) · Select Excel and/or CODE A / B / C"
        )
        ctk.CTkLabel(
            footer, textvariable=self.status_var, font=ui_font(11),
            text_color=COLOR_TEXT_MUTED, anchor="w",
        ).pack(side="left", fill="x", expand=True)
        self.ttl_led = ctk.CTkLabel(footer, text="● TTL", font=ui_font(10, "bold"), text_color=LED_IDLE)
        self.ttl_led.pack(side="right", padx=(0, 4))
        self.rx_led = ctk.CTkLabel(footer, text="● USB RX", font=ui_font(10, "bold"), text_color=LED_IDLE)
        self.rx_led.pack(side="right", padx=(8, 4))

    def _refresh_ports(self) -> None:
        ports = find_serial_ports() or [default_port()]
        self.port_combo.configure(values=ports)
        if self.port_combo.get() not in ports:
            self.port_combo.set(ports[0])
            self.port_var.set(ports[0])

    def _copy_code(self, key: str) -> None:
        text = str(self.code_vars.get(key, ctk.StringVar()).get() or "").strip()
        if not text or text == "—":
            return
        try:
            self.clipboard_clear()
            self.clipboard_append(text)
            self.status_var.set(f"Copied CODE {key}: {text}")
        except Exception:
            pass

    def _sync_meta_chips(self) -> None:
        raw = str(self.meta_var.get() or "")
        oe = supplier = freq = "—"
        # Formats: "OE: x  ·  Supplier: y  ·  Freq: z" or similar
        parts = [p.strip() for p in raw.replace("·", "|").split("|")]
        for part in parts:
            low = part.lower()
            if low.startswith("oe"):
                oe = part.split(":", 1)[-1].strip() or "—"
            elif "supplier" in low:
                supplier = part.split(":", 1)[-1].strip() or "—"
            elif "freq" in low:
                freq = part.split(":", 1)[-1].strip() or "—"
        if hasattr(self, "meta_chip_oe"):
            self.meta_chip_oe.configure(text=f"OE {oe}")
            self.meta_chip_supplier.configure(text=f"Supplier {supplier}")
            self.meta_chip_freq.configure(text=f"Freq {freq}")

    def _clear_telemetry(self, phase: str = "Waiting for Board read…") -> None:
        """Hard-reset tiles (session reset only). Prefer keeping last reading during a run."""
        if hasattr(self, "tel_id"):
            self.tel_id.set("—")
            self.tel_temp.set("—")
            self.tel_pressure.set("—")
            self.tel_rssi.set("— needs SDR" if not self._sdr_enabled() else "—")
            self.tel_batt.set("—")
            self.reading_var.set("Waiting for Board telemetry…")
        if hasattr(self, "tel_phase_var"):
            self.tel_phase_var.set(phase)
        self._last_board_tel = None

    @staticmethod
    def _tel_blank(val: str | None) -> bool:
        t = str(val or "").strip()
        return (not t) or t.lower() in {"na", "n/a", "—", "-", "none", "null"}

    def _set_telemetry(
        self,
        *,
        sensor_id: str | None = None,
        temperature: str | None = None,
        pressure: str | None = None,
        rssi: str | None = None,
        voltage: str | None = None,
        battery_percentage: str | None = None,
        phase: str | None = None,
        replace: bool = False,
        keep_last_on_empty: bool = True,
    ) -> None:
        """Push Board USB-TTL live values into the telemetry tiles."""
        if replace:
            self._clear_telemetry(phase or "Board reading…")

        sid = str(sensor_id or "").strip()
        temp = str(temperature or "").strip()
        press = str(pressure or "").strip()
        rs = str(rssi or "").strip()
        batt = _fmt_batt(voltage, battery_percentage)
        batt_disp = (
            batt.removeprefix("Batt ").strip()
            if batt and batt not in {"Batt —", "Batt na"}
            else ""
        )
        has_any = (
            not self._tel_blank(sid)
            or not self._tel_blank(temp)
            or not self._tel_blank(press)
            or not self._tel_blank(rs)
            or bool(batt_disp)
        )

        # No new values (e.g. trigger miss mid-row): keep last good Board reading visible.
        if not has_any:
            if keep_last_on_empty and getattr(self, "_last_board_tel", None):
                if phase and hasattr(self, "tel_phase_var"):
                    self.tel_phase_var.set(phase)
                return
            if phase and hasattr(self, "tel_phase_var"):
                self.tel_phase_var.set(phase or "No Board telemetry yet")
            return

        if not self._tel_blank(sid):
            self.tel_id.set(sid.upper() if len(sid) <= 16 else sid)
        if not self._tel_blank(temp):
            if "°" in temp or "C" in temp.upper():
                self.tel_temp.set(temp)
            else:
                try:
                    self.tel_temp.set(f"{float(temp):.0f} °C")
                except ValueError:
                    self.tel_temp.set(f"{temp} °C")
        if not self._tel_blank(press):
            # Compact tile text — avoid truncated "… bar / kPa" overflow.
            psi, bar = self._split_pressure(press)
            if psi != "na" and bar != "na":
                self.tel_pressure.set(f"{psi} PSI · {bar} bar")
            elif psi != "na":
                self.tel_pressure.set(f"{psi} PSI")
            else:
                self.tel_pressure.set(press.split("/")[0].strip() or press)
        if not self._tel_blank(rs):
            self.tel_rssi.set(rs)
        elif not self._sdr_enabled() and self.tel_rssi.get() in {"—", ""}:
            self.tel_rssi.set("— needs SDR")
        if batt_disp:
            self.tel_batt.set(batt_disp)

        self._last_board_tel = {
            "id": self.tel_id.get(),
            "temp": self.tel_temp.get(),
            "pressure": self.tel_pressure.get(),
            "rssi": self.tel_rssi.get(),
            "batt": self.tel_batt.get(),
        }
        if phase and hasattr(self, "tel_phase_var"):
            self.tel_phase_var.set(phase)
        elif hasattr(self, "tel_phase_var"):
            self.tel_phase_var.set("Board live · USB-TTL")

        self.reading_var.set(
            f"ID  {self.tel_id.get()}  ·  {self.tel_temp.get()}  ·  "
            f"{self.tel_pressure.get()}  ·  RSSI {self.tel_rssi.get()}  ·  {self.tel_batt.get()}"
        )

    @staticmethod
    def _split_pressure(raw: str | None) -> tuple[str, str]:
        """Return (PSI, bar) display strings from Board pressure text."""
        text = str(raw or "").strip()
        if not text or text.lower() in {"na", "n/a", "—", "-"}:
            return "na", "na"
        psi = "na"
        bar = "na"
        # Formats: "12.34 PSI / 0.850 bar / 85.00 kPa" or "12.34 PSI" or bare number
        import re

        m_psi = re.search(r"([-+]?\d+(?:\.\d+)?)\s*PSI", text, re.I)
        m_bar = re.search(r"([-+]?\d+(?:\.\d+)?)\s*bar", text, re.I)
        m_kpa = re.search(r"([-+]?\d+(?:\.\d+)?)\s*kPa", text, re.I)
        if m_psi:
            psi = f"{float(m_psi.group(1)):.2f}"
        if m_bar:
            bar = f"{float(m_bar.group(1)):.3f}"
        if psi == "na" and m_kpa:
            kpa = float(m_kpa.group(1))
            psi = f"{kpa * 0.1450377377:.2f}"
            bar = f"{kpa / 100.0:.3f}"
        if psi != "na" and bar == "na":
            bar = f"{float(psi) / 14.50377377:.3f}"
        if psi == "na":
            try:
                # Bare PSI number
                n = float(text.split()[0])
                psi = f"{n:.2f}"
                bar = f"{n / 14.50377377:.3f}"
            except (TypeError, ValueError, IndexError):
                pass
        return psi, bar

    def _board_match_values(
        self,
        *,
        excel_row: int,
        vehicle: str,
        codes: str,
        sensor_id: str,
        performance: str,
        temperature: str,
        voltage: str,
        pressure: str,
        rssi: str,
        reason: str,
        seen: str | None = None,
        count_read: bool = True,
    ) -> tuple:
        """Build one SDR-style match row for the Board history table."""
        sid = (sensor_id or "na").strip() or "na"
        psi, bar = self._split_pressure(pressure)
        temp = str(temperature or "na").strip() or "na"
        if temp.lower() not in {"na", "n/a"} and "°" not in temp:
            try:
                temp = f"{float(temp):.0f}"
            except ValueError:
                pass
        batt = str(voltage or "na").strip() or "na"
        rs = str(rssi or "").strip()
        if not rs or rs.lower() in {"na", "n/a", "—", "-"}:
            rs = "na"
        elif rs.isdigit():
            rs = f"board {rs}"
        key = sid.upper() if sid.lower() not in {"na", "n/a", ""} else f"row:{excel_row}"
        if not hasattr(self, "_board_read_counts"):
            self._board_read_counts = {}
        if count_read:
            self._board_read_counts[key] = int(self._board_read_counts.get(key, 0)) + 1
            reads = self._board_read_counts[key]
        else:
            reads = 1
        stamp = seen if seen is not None else datetime.now().strftime("%H:%M:%S")
        if not stamp:
            stamp = "—"
        return (
            excel_row,
            vehicle or "—",
            codes or "—",
            sid,
            (performance or "SKIP").upper(),
            psi,
            bar,
            temp,
            batt,
            rs,
            reads,
            stamp,
            reason or "",
        )

    def _log_board(self, message: str, level: str = "info") -> None:
        if hasattr(self, "activity_terminal"):
            try:
                self.activity_terminal.write(message, level=level)
            except Exception:
                pass

    def _apply_history_filter(self, *_args) -> None:
        """Rebuild the tree from full history using filter + search."""
        if not hasattr(self, "tree"):
            return
        filt = ""
        try:
            filt = str(self.history_filter.get() or "")
        except Exception:
            filt = "All results"
        q = ""
        try:
            q = str(self.history_search.get() or "").strip().lower()
        except Exception:
            q = ""
        rows = list(getattr(self, "_history_all", []) or [])
        if not rows:
            for iid in self.tree.get_children():
                vals = self.tree.item(iid, "values")
                tags = self.tree.item(iid, "tags")
                rows.append((vals, tags))
            self._history_all = list(rows)
        for iid in self.tree.get_children():
            self.tree.delete(iid)
        shown = 0
        for vals, tags in rows:
            # result column index 4 in SDR-style Board matches table
            result = str((vals[4] if len(vals) > 4 else "") or "").upper()
            if filt.startswith("Passed") and result != "OK":
                continue
            if filt.startswith("Failed") and result != "NOK":
                continue
            if q:
                hay = " ".join(str(v) for v in vals).lower()
                if q not in hay:
                    continue
            self.tree.insert("", "end", values=vals, tags=tags)
            shown += 1
        if hasattr(self, "table_count_label"):
            total = len(rows)
            label = (
                f"{shown} of {total} test{'s' if total != 1 else ''}"
                if shown != total
                else f"{total} test{'s' if total != 1 else ''}"
            )
            self.table_count_label.configure(text=label)

    def _refresh_ready_banner(self) -> None:
        if self.source_xlsx and self.source_xlsx.is_file():
            n = vehicle_row_count(self.source_xlsx)
            port = port_device(self.port_combo.get()) or port_device(default_port()) or "COM?"
            if n > 0:
                self.vehicle_var.set(
                    f"{self.source_xlsx.name} — {n:,} vehicles · Run chunk / Run full"
                )
                self.status_var.set(f"Ready on {port} — {n:,} catalog rows · waiting for Run")
            else:
                self.vehicle_var.set(f"{self.source_xlsx.name} loaded — no catalog rows found")
                self.status_var.set(f"Ready on {port} — check Excel or add custom codes")
        else:
            self.vehicle_var.set("Ready — choose Excel or enter custom codes")
            self.status_var.set("Choose an Excel database and/or enter custom CODE A / B / C")

    def _reset_display(self) -> None:
        self._update_stat_labels(pending=0, done=0, ok=0, nok=0, total=0)
        self.meta_var.set("OE: —   ·   Supplier: —   ·   Freq: —")
        self.row_counter.configure(text="ROW —")
        self._clear_telemetry("Ready — start a Board run")
        for key in self.code_vars:
            self.code_vars[key].set("—")
        self.progress.set(0)
        self.progress_label.configure(text="Progress: 0.0% (0 of 0)")
        self.speed_label.configure(text="")
        self.elapsed_label.configure(text="")
        if hasattr(self, "eta_label"):
            self.eta_label.configure(text="")
        self._history_all = []
        self._refresh_table_count()
        self._refresh_ready_banner()

    def _fit_tree_columns(self, _event=None) -> None:
        """Stretch live columns to the tree width so no horizontal scrollbar is needed."""
        if not hasattr(self, "tree"):
            return
        try:
            width = int(self.tree.winfo_width())
        except tk.TclError:
            return
        if width < 80:
            return
        # Leave a little room for the vertical scrollbar border.
        usable = max(200, width - 4)
        weights = getattr(self, "_tree_col_weights", None) or {}
        total_w = sum(weights.values()) or 1.0
        for col, weight in weights.items():
            col_w = max(36, int(usable * (weight / total_w)))
            self.tree.column(col, width=col_w)

    def _refresh_table_count(self) -> None:
        n = len(getattr(self, "_history_all", None) or self.tree.get_children())
        if hasattr(self, "table_count_label"):
            self.table_count_label.configure(text=f"{n} test{'s' if n != 1 else ''}")

    def _note_sdr_compare(self, sdr_compare: str) -> None:
        cmp = (sdr_compare or "").strip().upper()
        if cmp in {"", "PENDING", "NA", "N/A"}:
            return
        if cmp == "SUCCESS":
            self._sdr_agree += 1
        elif cmp == "FAIL":
            self._sdr_disagree += 1

    def _clear_readings(self) -> None:
        self._history_all = []
        self._board_read_counts = {}
        for item in self.tree.get_children():
            self.tree.delete(item)
        self._sdr_agree = 0
        self._sdr_disagree = 0
        self._session_rows.clear()
        self._refresh_table_count()
        self._log_board("Session cleared — previous Board readings cleared.", level="dim")
        if hasattr(self, "charts"):
            self.charts.reset()
        if hasattr(self, "trend_charts"):
            self.trend_charts.reset()

    def _record_session_row(self, event: ProgressEvent) -> None:
        perf = str(event.performance or "").strip().upper()
        if perf not in ("OK", "NOK"):
            return
        sid = str(event.sensor_id or "").strip()
        if sid.lower() in {"na", "n/a", "none", "-", "—"}:
            sid = ""
        self._session_rows.append(
            {
                "sensor_id": sid,
                "board_performance": perf,
                "nok_reason": str(event.reason or ""),
                "sdr_compare": str(event.sdr_compare or ""),
                "sdr_reason": str(event.sdr_reason or ""),
                "excel_row": int(event.excel_row or 0),
                "make": str(event.make or ""),
                "model": str(event.model or ""),
                "duration_s": event.duration_s,
            }
        )

    def get_tested_rows(self) -> list[dict]:
        """Board rows for comparative analysis — current session, else SQLite resume."""
        if self._session_rows:
            return list(self._session_rows)
        return self._hydrate_session_rows_from_db()

    def _hydrate_session_rows_from_db(self) -> list[dict]:
        """Reload Board OK/NOK rows from bench.sqlite so Compare survives a restart."""
        try:
            from tpms_bench.results_db import connect, fetch_all
            from tpms_bench.runner import DB_PATH
        except Exception:
            return []
        if not DB_PATH.exists():
            return []
        try:
            db = connect(DB_PATH)
            rows = fetch_all(db)
            db.close()
        except Exception:
            return []
        hydrated: list[dict] = []
        for row in rows:
            perf = str(row["board_performance"] or "").strip().upper()
            if perf not in ("OK", "NOK"):
                continue
            sid = str(row["sensor_id"] or "").strip()
            if sid.lower() in {"na", "n/a", "none", "-", "—"}:
                sid = ""
            hydrated.append(
                {
                    "sensor_id": sid,
                    "board_performance": perf,
                    "nok_reason": str(row["nok_reason"] or ""),
                    "sdr_compare": str(row["sdr_compare"] or ""),
                    "sdr_reason": str(row["sdr_reason"] or ""),
                    "excel_row": int(row["excel_row"] or 0),
                    "make": str(row["make"] or ""),
                    "model": str(row["model"] or ""),
                    "duration_s": row["duration_s"],
                    "_db_row": row,
                }
            )
        if hydrated and not self._session_rows:
            # Store without the private _db_row helper for Compare consumers.
            self._session_rows = [
                {k: v for k, v in item.items() if k != "_db_row"} for item in hydrated
            ]
        return list(self._session_rows) if self._session_rows else []

    def _restore_session_from_db(self) -> int:
        """Repopulate the live Board table from SQLite after an app restart."""
        try:
            from tpms_bench.results_db import connect, counts, fetch_all
            from tpms_bench.runner import DB_PATH
        except Exception:
            return 0
        if not DB_PATH.exists():
            return 0
        try:
            db = connect(DB_PATH)
            rows = fetch_all(db)
            stats = counts(db)
            db.close()
        except Exception:
            return 0
        if not rows:
            return 0

        self._session_rows.clear()
        self._history_all = []
        for item in self.tree.get_children():
            self.tree.delete(item)

        shown = 0
        for row in rows:
            perf = str(row["board_performance"] or "").strip().upper()
            if perf not in ("OK", "NOK", "SKIP"):
                continue
            sid = str(row["sensor_id"] or "").strip()
            if sid.lower() in {"na", "n/a", "none", "-", "—"}:
                sid_disp = "na"
                sid = ""
            else:
                sid_disp = sid
            make = str(row["make"] or "")
            model = str(row["model"] or "")
            code_a = str(row["code_a"] or "")
            code_b = str(row["code_b"] or "")
            code_c = str(row["code_c"] or "")
            temp = str(row["temperature"] or "na")
            volt = _fmt_batt(
                str(row["battery_voltage"] or ""),
                str(row["battery_percentage"] or ""),
            ).removeprefix("Batt ").strip() or "na"
            press = str(row["pressure"] if "pressure" in row.keys() else "na") or "na"
            rssi = str(row["rssi"] if "rssi" in row.keys() else "na") or "na"
            reason = str(row["nok_reason"] or "")
            excel_row = int(row["excel_row"] or 0)
            vals = self._board_match_values(
                excel_row=excel_row,
                vehicle=f"{make} {model}".strip(),
                codes=f"{code_a}  {code_b}  {code_c}".strip(),
                sensor_id=sid_disp,
                performance=perf,
                temperature=temp,
                voltage=volt,
                pressure=press,
                rssi=rssi,
                reason=reason,
                seen="—",
                count_read=False,
            )
            tags = (perf,)
            self._history_all.insert(0, (vals, tags))
            self.tree.insert("", 0, values=vals, tags=tags)
            shown += 1
            if perf in ("OK", "NOK"):
                self._session_rows.append(
                    {
                        "sensor_id": sid,
                        "board_performance": perf,
                        "nok_reason": reason,
                        "sdr_compare": str(row["sdr_compare"] or ""),
                        "sdr_reason": str(row["sdr_reason"] or ""),
                        "excel_row": excel_row,
                        "make": make,
                        "model": model,
                        "duration_s": row["duration_s"],
                    }
                )
            self._note_sdr_compare(str(row["sdr_compare"] or ""))

        done = int(stats.get("done") or shown)
        ok = int(stats.get("OK") or 0)
        nok = int(stats.get("NOK") or 0)
        self._update_stat_labels(pending=0, done=done, ok=ok, nok=nok, total=done)
        self._refresh_table_count()
        # Seed live tiles from the newest row that has a real sensor ID.
        # columns: num,vehicle,codes,id,result,psi,bar,temp,batt,rssi,reads,time,reason
        for vals, _tags in list(getattr(self, "_history_all", []) or []):
            sid = str(vals[3] if len(vals) > 3 else "")
            if self._tel_blank(sid):
                continue
            psi = str(vals[5] if len(vals) > 5 else "")
            self._set_telemetry(
                sensor_id=sid,
                temperature=str(vals[7] if len(vals) > 7 else ""),
                pressure=f"{psi} PSI" if psi and psi != "na" else "",
                rssi=str(vals[9] if len(vals) > 9 else ""),
                voltage=str(vals[8] if len(vals) > 8 else ""),
                phase="Restored · last Board reading",
            )
            break
        if hasattr(self, "charts"):
            history = [str(r.get("board_performance") or "") for r in self._session_rows]
            self.charts.set_counts(ok, nok, history)
        if hasattr(self, "trend_charts"):
            history = [str(r.get("board_performance") or "") for r in self._session_rows]
            volts: list[float] = []
            temps: list[float] = []
            for row in rows:
                try:
                    v = str(row["battery_voltage"] or "").strip()
                    if v and v.lower() not in {"na", "n/a", "—", "-", "ok", "low"}:
                        volts.append(float(v))
                except (TypeError, ValueError):
                    pass
                try:
                    t = str(row["temperature"] or "").strip()
                    if t and t.lower() not in {"na", "n/a", "—", "-"}:
                        temps.append(float(t))
                except (TypeError, ValueError):
                    pass
            self.trend_charts.set_from_history(results=history, voltages=volts, temps=temps)
        return shown

    def focus_code_fields(self) -> None:
        try:
            self.manual_entries["A"].focus_set()
        except Exception:
            pass

    def _entry_text(self, entry: ctk.CTkEntry) -> str:
        value = (entry.get() or "").strip()
        try:
            placeholder = str(entry.cget("placeholder_text") or "")
        except Exception:
            placeholder = ""
        if placeholder and value == placeholder:
            return ""
        return value

    def _read_manual_fields(self, *, require_complete: bool) -> ManualCode | None:
        values = {key: self._entry_text(self.manual_entries[key]) for key in ("A", "B", "C")}
        filled = [v for v in values.values() if v]
        if not filled:
            return None
        if len(filled) != 3:
            if require_complete:
                messagebox.showinfo(
                    "Custom opcodes",
                    "Enter CODE A, CODE B and CODE C together before adding.",
                    parent=self._toplevel(),
                )
                raise ValueError("incomplete custom codes")
            return None
        for key, raw in values.items():
            if parse_code(raw) is None:
                messagebox.showinfo("Custom opcodes", f"CODE {key} is not a valid hex opcode: {raw}", parent=self._toplevel())
                raise ValueError("invalid custom code")
        label = self._entry_text(self.manual_label_entry) or f"Manual code {len(self.custom_codes) + 1}"
        make = self._entry_text(self.manual_make_entry)
        model = self._entry_text(self.manual_model_entry)
        freq = self._entry_text(self.manual_freq_entry) or "433.92"
        return ManualCode(
            code_a=values["A"].upper(),
            code_b=values["B"].upper(),
            code_c=values["C"].upper(),
            label=label,
            make=make,
            model=model,
            freq=freq,
        )

    def add_custom_code(self) -> None:
        try:
            extra = self._read_manual_fields(require_complete=True)
        except ValueError:
            return
        if extra is None:
            messagebox.showinfo("Custom opcodes", "Type CODE A, CODE B and CODE C, then click Add Code.", parent=self._toplevel())
            return
        self.custom_codes.append(extra)
        for entry in self.manual_entries.values():
            entry.delete(0, "end")
        for entry in (
            self.manual_label_entry,
            self.manual_make_entry,
            self.manual_model_entry,
            self.manual_freq_entry,
        ):
            entry.delete(0, "end")
        self._refresh_codes_list()
        self.status_var.set(f"Added custom code {extra.label}: {extra.code_a} / {extra.code_b} / {extra.code_c}")
        self.focus_code_fields()

    def _remove_custom_code(self, index: int) -> None:
        if 0 <= index < len(self.custom_codes):
            removed = self.custom_codes.pop(index)
            self._refresh_codes_list()
            self.status_var.set(f"Removed custom code {removed.label}")

    def _refresh_codes_list(self) -> None:
        for child in self.codes_list.winfo_children():
            child.destroy()
        if not self.custom_codes:
            self.codes_hint = ctk.CTkLabel(
                self.codes_list,
                text="No custom codes added yet.",
                font=ctk.CTkFont(size=10),
                text_color=COLOR_TEXT_MUTED,
                anchor="w",
            )
            self.codes_hint.pack(anchor="w")
            return
        for index, code in enumerate(self.custom_codes):
            row = ctk.CTkFrame(self.codes_list, fg_color=COLOR_CYAN_BG, corner_radius=6)
            row.pack(side="left", padx=(0, 6), pady=2)
            meta = " · ".join(
                part
                for part in (
                    code.make,
                    code.model,
                    code.freq,
                )
                if part
            )
            text = f"{code.label}: {code.code_a}/{code.code_b}/{code.code_c}"
            if meta:
                text = f"{text}  ({meta})"
            ctk.CTkLabel(
                row,
                text=text,
                font=ctk.CTkFont(family="Consolas", size=11),
                text_color=COLOR_TEXT,
            ).pack(side="left", padx=(8, 4), pady=3)
            ctk.CTkButton(
                row,
                text="×",
                width=28,
                height=22,
                font=ctk.CTkFont(size=12),
                fg_color=COLOR_BTN_SECONDARY,
                hover_color=COLOR_BTN_SECONDARY_HOVER,
                command=lambda i=index: self._remove_custom_code(i),
            ).pack(side="right", padx=(0, 4), pady=2)

    def select_excel(self) -> None:
        path = ask_open_path(
            "Select TPMS Board Excel database (opcodes)",
            [("Excel files", "*.xlsx"), ("All files", "*.*")],
            parent=self._toplevel(),
        )
        if not path:
            return
        self._load_excel(Path(path))

    def remove_excel(self) -> None:
        """Clear the selected Excel catalog (custom CODE A/B/C stay)."""
        if self._state_label in ("RUNNING", "PAUSED"):
            messagebox.showinfo(
                "Remove Excel",
                "Stop the Board test before removing the Excel file.",
                parent=self._toplevel(),
            )
            return
        if not self.source_xlsx:
            self._update_excel_buttons()
            return
        name = self.source_xlsx.name
        if not messagebox.askyesno(
            "Remove Excel",
            f"Remove “{name}” from this session?\n\n"
            "Custom CODE A / B / C entries are kept. You can choose another Excel later.",
            parent=self._toplevel(),
        ):
            return
        self.source_xlsx = None
        self.file_label.configure(text="No file selected", text_color=COLOR_ORANGE)
        self._update_excel_buttons()
        self._refresh_ready_banner()
        self.status_var.set(f"Removed {name} — choose another Excel or use custom codes")

    def _update_excel_buttons(self) -> None:
        has_file = bool(self.source_xlsx and Path(self.source_xlsx).exists())
        if hasattr(self, "remove_excel_btn"):
            self.remove_excel_btn.configure(state="normal" if has_file else "disabled")

    def _load_excel(self, path: Path) -> None:
        self.source_xlsx = path
        self.file_label.configure(text=self.source_xlsx.name, text_color=COLOR_GREEN)
        self._update_excel_buttons()
        self._refresh_ready_banner()

    def _autoload_default_excel(self) -> None:
        """Preselect Hamaton catalog (or FYRQOM_SOURCE_XLSX override) so a run can start."""
        if self.source_xlsx:
            return
        override = os.environ.get("FYRQOM_SOURCE_XLSX", "").strip()
        if override:
            path = Path(override).expanduser()
            if not path.is_absolute():
                path = app_root() / path
            if path.is_file():
                self._load_excel(path)
                return
        if DEFAULT_HAMATON_XLSX.is_file():
            self._load_excel(DEFAULT_HAMATON_XLSX)

    def previous_session_count(self) -> int:
        """How many Board OK/NOK rows are stored from the last session."""
        try:
            from tpms_bench.results_db import connect, fetch_all
            from tpms_bench.runner import DB_PATH
        except Exception:
            return 0
        if not DB_PATH.exists():
            return 0
        try:
            db = connect(DB_PATH)
            rows = fetch_all(db)
            db.close()
        except Exception:
            return 0
        return sum(1 for row in rows if str(row["board_performance"] or "").strip().upper() in {"OK", "NOK"})

    def restore_previous_session(self) -> int:
        """Hydrate the Board table from SQLite. Returns restored row count."""
        restored = self._restore_session_from_db()
        if restored:
            self.status_var.set(f"Restored {restored} Board result row(s) from last session")
            self._set_state_badge("IDLE")
        return restored

    def clear_session_silent(self) -> None:
        """Clear Board results without a confirmation dialog (startup New Session)."""
        try:
            reset_session_db()
        except Exception:
            pass
        if self.source_xlsx and self.source_xlsx.is_file():
            try:
                from tpms_bench.excel_io import copy_workbook
                from tpms_bench.runner import OUT_XLSX

                copy_workbook(self.source_xlsx, OUT_XLSX, force=True)
            except Exception:
                pass
        self._clear_readings()
        self._reset_display()
        self._set_state_badge("IDLE")
        self.status_var.set("New session — previous Board results cleared")

    def reset_session(self) -> None:
        if self.worker and self.worker.is_alive():
            messagebox.showwarning("Session Active", "Stop the current test before resetting.", parent=self._toplevel())
            return
        if not messagebox.askyesno("Reset Session", "Clear all TPMS Board results and reset counters?", parent=self._toplevel()):
            return
        reset_session_db()
        if self.source_xlsx and self.source_xlsx.is_file():
            try:
                from tpms_bench.excel_io import copy_workbook
                from tpms_bench.runner import OUT_XLSX

                copy_workbook(self.source_xlsx, OUT_XLSX, force=True)
            except Exception:
                pass
        self._clear_readings()
        self._reset_display()
        self._set_state_badge("IDLE")

    def selected_chunk_size(self) -> int:
        try:
            return max(1, int(self.chunk_var.get()))
        except (TypeError, ValueError):
            return 100

    def start_test(self, *, skip_sdr: bool = True, chunk_size: int | None = None, run_full: bool = False) -> bool:
        """Start the Board bench. Returns False if start was blocked (e.g. no Excel/codes)."""
        if self.worker and self.worker.is_alive():
            if self.is_paused and self.runner:
                self.toggle_pause()
            return True

        extra_codes = list(self.custom_codes)
        try:
            typed = self._read_manual_fields(require_complete=bool(self._entry_text(self.manual_entries["A"]) or self._entry_text(self.manual_entries["B"]) or self._entry_text(self.manual_entries["C"])))
        except ValueError:
            return False
        if typed:
            extra_codes.append(typed)

        if (not self.source_xlsx or not self.source_xlsx.exists()) and not extra_codes:
            messagebox.showinfo(
                "Select Excel or add codes",
                "Choose an Excel database and/or add custom CODE A / B / C before starting.",
                parent=self._toplevel(),
            )
            return False

        # Board-only (USB-TTL): start immediately — SDR is optional and never blocks.
        if skip_sdr:
            return self._start_test_now(
                skip_sdr=True,
                chunk_size=chunk_size,
                run_full=run_full,
                extra_codes=extra_codes,
            )

        freq_mhz = self._preferred_sdr_freq_mhz(extra_codes)
        # Board+SDR mode only: arm rtl_433 before ABC when a dongle is present.
        if callable(self.on_prepare_sdr):
            try:
                ready = bool(self.on_prepare_sdr(freq_mhz=freq_mhz))
            except Exception:
                ready = False
            if not ready:
                self.status_var.set(
                    f"Starting SDR on {(freq_mhz if freq_mhz is not None else 433.92):.2f} MHz first — Board starts in a few seconds…"
                )
                self.start_btn.configure(state="disabled")
                if hasattr(self, "full_btn"):
                    self.full_btn.configure(state="disabled")
                self._set_state_badge("RUNNING")
                self.after(
                    4500,
                    lambda: self._start_test_now(
                        skip_sdr=False,
                        chunk_size=chunk_size,
                        run_full=run_full,
                        extra_codes=extra_codes,
                    ),
                )
                return True

        return self._start_test_now(
            skip_sdr=False,
            chunk_size=chunk_size,
            run_full=run_full,
            extra_codes=extra_codes,
        )

    def _preferred_sdr_freq_mhz(self, extra_codes: list | None = None) -> float | None:
        """Best center frequency for upcoming Board rows (custom Freq, else Excel)."""
        codes = list(extra_codes or self.custom_codes or [])
        for code in codes:
            text = str(getattr(code, "freq", "") or "").strip()
            if not text:
                continue
            try:
                from tpms_bench.rtl433 import parse_freq_hz

                return parse_freq_hz(text) / 1_000_000.0
            except Exception:
                try:
                    return float(text.replace("MHz", "").replace("M", "").strip())
                except ValueError:
                    pass
        if codes:
            return 433.92
        # Excel-only run: peek first catalog Freq so SDR is armed on the right band
        # before ABC (avoids retune race that NOKs Hamaton DB rows).
        path = self.source_xlsx
        if path and Path(path).is_file():
            try:
                from openpyxl import load_workbook
                from tpms_bench.rtl433 import parse_freq_hz

                wb = load_workbook(path, read_only=True, data_only=True)
                ws = wb[wb.sheetnames[0]]
                headers = {
                    str(c.value).strip(): int(c.column)
                    for c in ws[1]
                    if c.value not in (None, "")
                }
                freq_col = headers.get("Freq") or headers.get("Frequency") or 10
                code_a_col = headers.get("CODEA") or headers.get("Code A") or 16
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row:
                        continue
                    vals = list(row)
                    ca = vals[code_a_col - 1] if len(vals) >= code_a_col else None
                    if ca in (None, ""):
                        continue
                    freq_val = vals[freq_col - 1] if len(vals) >= freq_col else None
                    if freq_val in (None, ""):
                        continue
                    mhz = parse_freq_hz(freq_val) / 1_000_000.0
                    wb.close()
                    return mhz
                wb.close()
            except Exception:
                pass
            return 433.92
        return None

    def _start_test_now(
        self,
        *,
        skip_sdr: bool = True,
        chunk_size: int | None = None,
        run_full: bool = False,
        extra_codes: list | None = None,
    ) -> bool:
        if self.worker and self.worker.is_alive():
            return True

        codes = list(extra_codes) if extra_codes is not None else list(self.custom_codes)
        # Keep existing Board results — only "Reset Session" clears. Resume unfinished rows.
        if not self._session_rows:
            self._hydrate_session_rows_from_db()
        self.is_paused = False
        self.start_time = time.monotonic()
        self.start_btn.configure(state="disabled")
        if hasattr(self, "full_btn"):
            self.full_btn.configure(state="disabled")
        self.pause_btn.configure(state="normal", text="Pause test")
        self.stop_btn.configure(state="normal")
        self._set_state_badge("RUNNING")

        size = 0 if run_full else (chunk_size if chunk_size is not None else self.selected_chunk_size())
        port = port_device(self.port_combo.get()) or port_device(default_port())
        self._session_skip_sdr = bool(skip_sdr)
        # Board-only: do not wire live SDR callbacks (otherwise the runner waits on rtl_433).
        self.runner = BenchRunner(
            port=port,
            source_xlsx=self.source_xlsx,
            resume=True,
            skip_sdr=skip_sdr,
            sdr_timeout=float(__import__("os").environ.get("FYRQOM_SDR_TIMEOUT", "10")),
            extra_codes=codes,
            on_progress=self.queue.put,
            chunk_size=size,
            retest_from=None,
            require_agree=False,
            agree_retries=1,
            live_sdr_get=None if skip_sdr else self.live_sdr_get,
            live_sdr_prepare=None if skip_sdr else self.live_sdr_prepare,
        )
        if skip_sdr:
            self.status_var.set(
                f"Board-only · USB-TTL on {port} · no SDR required · programming ABC…"
            )
        else:
            self.status_var.set(
                f"Board + SDR · USB-TTL on {port} · live rtl_433 compare enabled…"
            )
        self.worker = threading.Thread(target=self._run_safe, daemon=True, name="tpms-board")
        self.worker.start()
        return True

    def toggle_pause(self) -> None:
        if not self.runner or not self.worker or not self.worker.is_alive():
            return
        if not self.is_paused:
            self.runner.pause()
            self.is_paused = True
            self.pause_btn.configure(text="Resume")
            self._set_state_badge("PAUSED")
        else:
            self.runner.resume_run()
            self.is_paused = False
            self.pause_btn.configure(text="Pause test")
            self._set_state_badge("RUNNING")

    def stop_test(self) -> None:
        if self.runner:
            self.runner.request_stop()
        self.stop_btn.configure(state="disabled")
        self.pause_btn.configure(state="disabled", text="Pause test")
        self.is_paused = False
        self._set_state_badge("STOPPED")
        self.status_var.set("Stopping — aborting COM wait…")
        # If the worker is stuck opening/using a bad COM port, force UI recovery soon.
        self.after(800, self._ensure_stopped_ui)

    def _ensure_stopped_ui(self) -> None:
        """Recover controls if stop was requested but the worker never finished cleanly."""
        if self.worker and self.worker.is_alive():
            self.after(400, self._ensure_stopped_ui)
            return
        if self.start_btn.cget("state") == "disabled":
            self.start_btn.configure(state="normal")
            if hasattr(self, "full_btn"):
                self.full_btn.configure(state="normal")
            self.pause_btn.configure(state="disabled", text="Pause test")
            self.stop_btn.configure(state="disabled")
            self.is_paused = False
            if self._state_label in ("RUNNING", "PAUSED"):
                self._set_state_badge("STOPPED")
            self.status_var.set("Stopped")

    def _set_state_badge(self, text: str) -> None:
        self._state_label = text
        fg = {
            "RUNNING": COLOR_PILL_OK_FG,
            "PAUSED": COLOR_ORANGE,
            "STOPPED": COLOR_TEXT_DIM,
            "COMPLETED": COLOR_BLUE,
            "IDLE": COLOR_TEXT_MUTED,
            "ERROR": COLOR_RED,
        }
        bg = {
            "RUNNING": COLOR_PILL_OK_BG,
            "PAUSED": COLOR_ORANGE_BG,
            "STOPPED": COLOR_BG_PANEL,
            "COMPLETED": COLOR_BLUE_BG,
            "IDLE": COLOR_BG_PANEL,
            "ERROR": COLOR_RED_BG,
        }
        self.state_badge.configure(
            text=text,
            text_color=fg.get(text, COLOR_TEXT_MUTED),
            fg_color=bg.get(text, COLOR_BG_PANEL),
        )
        self._notify_status()

    def _notify_status(self) -> None:
        running = self._state_label in ("RUNNING", "PAUSED")
        try:
            self.on_status_change(running, self._state_label)
        except Exception:
            pass

    def is_running(self) -> bool:
        return bool(self.worker and self.worker.is_alive())

    def _push_tpms_range(self, event: ProgressEvent) -> None:
        # Soft-fill voltage/temp into trend hist without counting a result (row_done does that).
        if hasattr(self, "trend_charts"):
            volt = None
            temp = None
            try:
                v_raw = str(getattr(event, "voltage", "") or "").strip()
                if v_raw and v_raw.lower() not in {"na", "n/a", "—", "-"}:
                    volt = float(v_raw)
            except (TypeError, ValueError):
                volt = None
            try:
                t_raw = str(getattr(event, "temperature", "") or "").strip()
                if t_raw and t_raw.lower() not in {"na", "n/a", "—", "-"}:
                    temp = float(t_raw)
            except (TypeError, ValueError):
                temp = None
            if volt is None and temp is None:
                return
            try:
                self.trend_charts.push(voltage=volt, temp=temp)
            except Exception:
                pass

    def _poll_live_sdr(self) -> None:
        """Show SDR packets on the telemetry bar so a silent Board still looks alive."""
        try:
            snap = (
                self.live_sdr_get()
                if self._sdr_enabled() and self.live_sdr_get
                else None
            )
        except Exception:
            snap = None
        if snap:
            latest = None
            for reading in snap.values():
                ts = getattr(reading, "timestamp", None)
                if latest is None or (ts and getattr(latest, "timestamp", None) and ts > latest.timestamp):
                    latest = reading
            if latest is None:
                latest = next(iter(snap.values()), None)
            if latest is not None:
                sid = getattr(latest, "sensor_id", None) or "—"
                temp = getattr(latest, "display_temp", None) or "—"
                psi = getattr(latest, "display_pressure", None) or "—"
                rssi = getattr(latest, "display_rssi", None) or "—"
                self._set_telemetry(
                    sensor_id=str(sid),
                    temperature=str(temp),
                    pressure=str(psi),
                    rssi=str(rssi),
                )
        try:
            delay = 400 if self.winfo_ismapped() else 1200
        except tk.TclError:
            delay = 1200
        self.after(delay, self._poll_live_sdr)

    def _animate_badge(self) -> None:
        if self._state_label == "RUNNING":
            self._anim_step = (self._anim_step + 1) % 20
            pulse = abs(10 - (self._anim_step % 20))
            green = "#1A8A6C" if pulse > 4 else "#3EC9A5"
            self.state_badge.configure(text_color=green)
            if hasattr(self, "progress"):
                try:
                    self.progress.configure(progress_color=green)
                except Exception:
                    pass
        self.after(90, self._animate_badge)

    def _run_safe(self) -> None:
        try:
            assert self.runner is not None
            self.runner.run()
        except Exception as error:
            message = str(error)
            low = message.lower()
            if "permission denied" in low or "[errno 13]" in low:
                message = (
                    "Cannot write the Board results Excel — it is open in another program "
                    "(usually Microsoft Excel).\n\n"
                    "Close that file, then press START again.\n\n"
                    f"({error})"
                )
            self.queue.put(ProgressEvent(kind="error", message=message))

    def _open_file(self, path: Path) -> None:
        if not path.exists():
            messagebox.showinfo("File Not Found", str(path), parent=self._toplevel())
            return
        if platform.system() == "Windows":
            os.startfile(str(path))  # type: ignore[attr-defined]
        elif platform.system() == "Darwin":
            subprocess.run(["open", str(path)], check=False)
        else:
            subprocess.run(["xdg-open", str(path)], check=False)

    def export_excel(self) -> None:
        out = get_out_xlsx() if get_out_xlsx().exists() else OUT_XLSX
        if not out.exists():
            messagebox.showinfo("TPMS Board Report (Excel)", "No TPMS Board results yet. Run a test first.", parent=self._toplevel())
            return
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        dest = ask_save_path(
            "Save TPMS Board Report (Excel) — Hamaton bench, not SDR receiver",
            f"TPMS_Board_Report_{stamp}.xlsx",
            [("Excel — TPMS Board Report", "*.xlsx"), ("All files", "*.*")],
            ".xlsx",
            parent=self._toplevel(),
        )
        if not dest:
            return
        try:
            stamp_board_report_info(out)
            shutil.copy2(out, dest)
            self.status_var.set(f"TPMS Board Report (Excel) saved: {Path(dest).name}")
            messagebox.showinfo(
                "TPMS Board Report (Excel)",
                f"TPMS Board Report (Excel) saved.\n\nThis file is Hamaton board validation, not an SDR receiver report.\n\n{dest}",
                parent=self._toplevel(),
            )
            self._open_file(Path(dest))
        except Exception as error:
            messagebox.showerror("TPMS Board Report (Excel)", str(error), parent=self._toplevel())

    def export_pdf(self) -> None:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        dest = ask_save_path(
            "Save TPMS Board Report (PDF) — Hamaton bench, not SDR receiver",
            f"TPMS_Board_Report_{stamp}.pdf",
            [("PDF — TPMS Board Report", "*.pdf"), ("All files", "*.*")],
            ".pdf",
            parent=self._toplevel(),
        )
        if not dest:
            return
        try:
            saved = build_pdf(Path(dest))
            self.status_var.set(f"TPMS Board Report (PDF) saved: {saved.name}")
            messagebox.showinfo(
                "TPMS Board Report (PDF)",
                f"TPMS Board Report (PDF) saved.\n\nThis file is Hamaton board validation, not an SDR receiver report.\n\n{saved}",
                parent=self._toplevel(),
            )
            self._open_file(saved)
        except Exception as error:
            messagebox.showerror("TPMS Board Report (PDF)", str(error), parent=self._toplevel())

    def _drain(self) -> None:
        # Cap work per tick so SDR UI timers still run while the Board is busy.
        try:
            for _ in range(40):
                self._apply(self.queue.get_nowait())
        except queue.Empty:
            pass
        self.after(80, self._drain)

    def _flash_row(self, item_id: str, tag: str, step: int = 0) -> None:
        # Skip animation under load — flashing every row freezes the UI (and SDR).
        if step == 0 and len(self.queue.queue) > 5:
            try:
                if self.tree.exists(item_id):
                    self.tree.item(item_id, tags=(tag,))
            except tk.TclError:
                pass
            return
        if not self.tree.exists(item_id):
            return
        try:
            if step < 2:
                self.tree.item(item_id, tags=("flash",))
                self.after(50, lambda i=item_id, t=tag, s=step + 1: self._flash_row(i, t, s))
            else:
                self.tree.item(item_id, tags=(tag,))
        except tk.TclError:
            pass

    def _end_session(self, label: str, message: str) -> None:
        self.start_btn.configure(state="normal")
        if hasattr(self, "full_btn"):
            self.full_btn.configure(state="normal")
        self.pause_btn.configure(state="disabled", text="Pause test")
        self.stop_btn.configure(state="disabled")
        self.is_paused = False
        self._set_state_badge(label)
        try:
            out = get_out_xlsx()
            if out.exists():
                stamp_board_report_info(out)
            build_pdf()
        except Exception:
            pass
        self.status_var.set(message)
        # Only stop SDR when this session actually used it.
        if self._sdr_enabled() and callable(self.on_sdr_release):
            try:
                self.on_sdr_release(reason=label, message=message or "")
            except Exception:
                pass

    def _release_sdr_after_chunk(self, message: str = "") -> None:
        if self._sdr_enabled() and callable(self.on_sdr_release):
            try:
                self.on_sdr_release(reason="CHUNK", message=message or "Chunk complete")
            except Exception:
                pass

    def _blink_led(self, widget: ctk.CTkLabel, active_color: str) -> None:
        widget.configure(text_color=active_color)
        self.after(350, lambda: widget.configure(text_color=LED_IDLE))

    def _set_board_version_ui(self, extras: dict, message: str = "") -> None:
        """Show firmware (TPMS software) first; other Chapter 4.0 fields secondary."""
        if not hasattr(self, "version_var"):
            return

        def _num(key: str):
            raw = extras.get(key)
            if raw in (None, ""):
                return None
            try:
                return int(raw)
            except (TypeError, ValueError):
                return raw

        if not extras:
            self.version_var.set(message or "Firmware — no reply")
            return

        fw = _num("firmware_version")
        if fw is None:
            fw = _num("software_version")
        if fw is None:
            self.version_var.set(message or "Firmware — no reply")
            return

        try:
            fw_text = f"{int(fw)}  (0x{int(fw):08X})"
        except (TypeError, ValueError):
            fw_text = str(fw)

        extras_parts = []
        for key, short in (
            ("hardware_version", "HW"),
            ("boot_version", "Boot"),
            ("trigger_database_version", "Trigger"),
            ("programming_database_version", "Prog DB"),
        ):
            val = _num(key)
            if val is None:
                continue
            try:
                extras_parts.append(f"{short} {int(val)}")
            except (TypeError, ValueError):
                extras_parts.append(f"{short} {val}")
        suffix = f"  ·  {' · '.join(extras_parts)}" if extras_parts else ""
        self.version_var.set(f"{fw_text}{suffix}")

    def query_firmware_version(self) -> None:
        """Open UART, run Chapter 4.0 Query Version (0x40), show firmware on the strip."""
        if self.worker and self.worker.is_alive():
            self.status_var.set("Stop the Board run before querying firmware")
            return
        if getattr(self, "_fw_query_busy", False):
            return
        self._fw_query_busy = True
        if hasattr(self, "query_fw_btn"):
            self.query_fw_btn.configure(state="disabled")
        self.version_var.set("Querying firmware…")
        self.status_var.set("Query Version (0x40) — reading firmware…")

        def _work() -> None:
            extras: dict = {}
            err = ""
            try:
                from tpms_bench.board import BoardSession
                from tpms_bench.uart import open_uart

                port = None
                try:
                    port = port_device(self.port_combo.get()) or None
                except Exception:
                    port = None
                ser = open_uart(port)
                try:
                    session = BoardSession(
                        ser,
                        on_comm=lambda ch, detail: self.after(
                            0, lambda: self._apply(
                                ProgressEvent(kind="comm", path=ch, message=detail)
                            )
                        ),
                    )
                    session.detect_transport()
                    version = session.query_board_version()
                    if version is not None:
                        extras = {
                            "firmware_version": int(version.software_version),
                            "hardware_version": int(version.hardware_version),
                            "boot_version": int(version.boot_version),
                            "software_version": int(version.software_version),
                            "trigger_database_version": int(version.trigger_database_version),
                            "programming_database_version": int(
                                version.programming_database_version
                            ),
                        }
                    else:
                        err = "no reply"
                finally:
                    try:
                        ser.close()
                    except Exception:
                        pass
            except Exception as exc:
                err = str(exc) or "query failed"

            def _done() -> None:
                self._fw_query_busy = False
                if hasattr(self, "query_fw_btn"):
                    self.query_fw_btn.configure(state="normal")
                if extras:
                    self._set_board_version_ui(extras)
                    fw = extras.get("firmware_version", extras.get("software_version"))
                    self.status_var.set(f"Firmware {fw}")
                    try:
                        self.on_process_update(
                            kind="board_version",
                            message=f"Firmware {fw}",
                            extras=dict(extras),
                        )
                    except Exception:
                        pass
                else:
                    self.version_var.set(f"Firmware — {err or 'no reply'}")
                    self.status_var.set(f"Query firmware failed · {err or 'no reply'}")

            self.after(0, _done)

        threading.Thread(target=_work, daemon=True).start()

    def _apply(self, event: ProgressEvent) -> None:
        try:
            self.on_process_update(
                kind=event.kind,
                path=str(event.path or ""),
                message=str(event.message or ""),
                excel_row=event.excel_row,
                sensor_id=str(event.sensor_id or ""),
                performance=str(event.performance or ""),
                code_a=str(event.code_a or ""),
                code_b=str(event.code_b or ""),
                code_c=str(event.code_c or ""),
                make=str(event.make or ""),
                model=str(event.model or ""),
                transport=str(event.transport or ""),
                tool_path=str(event.path or ""),
                temperature=str(event.temperature or ""),
                pressure=str(event.pressure or ""),
                voltage=str(event.voltage or ""),
                freq=str(event.freq or ""),
                rtl433_decoder=str(event.rtl433_decoder or ""),
                sdr_compare=str(event.sdr_compare or ""),
                extras=dict(getattr(event, "extras", None) or {}),
            )
        except Exception:
            pass

        if event.kind == "board_version":
            self._set_board_version_ui(getattr(event, "extras", None) or {}, event.message or "")
            return

        if event.kind == "comm":
            if event.path == "ttl":
                self._blink_led(self.ttl_led, LED_TTL_ON)
            elif event.path == "uart":
                self.rx_led.configure(text="● USB RX")
                self._blink_led(self.rx_led, LED_TTL_ON)
            elif event.message:
                self.status_var.set(event.message)
            return

        if event.kind == "phase":
            phase = str(event.path or "").lower()
            label = {
                "program": "Programming ABC…",
                "trigger": "LF trigger — waiting for sensor…",
                "query": "Querying Board telemetry…",
            }.get(phase, event.message or phase)
            if event.message:
                self.status_var.set(event.message)
            else:
                self.status_var.set(label)
            if hasattr(self, "tel_phase_var"):
                if getattr(self, "_last_board_tel", None) and phase in {"program", "trigger"}:
                    self.tel_phase_var.set(f"{label} · last reading held")
                else:
                    self.tel_phase_var.set(label)
            self._log_board(label, level="dim")
            return

        if event.kind == "started":
            extras = getattr(event, "extras", None) or {}
            if extras:
                self._set_board_version_ui(extras, event.message or "")
            self._total_rows = event.total
            self._set_state_badge("RUNNING")
            if event.transport:
                self.transport_var.set(event.transport)
                t = event.transport.upper()
                if "USB" in t:
                    self.rx_led.configure(text="● USB RX")
                else:
                    self.rx_led.configure(text="● USB RX")
            self._update_stat_labels(event.pending, event.done, event.ok, event.nok, event.total)
            self.status_var.set(event.message)

        elif event.kind == "row_start":
            self.vehicle_var.set(f"{event.make} {event.model} ({event.year})" if event.year else f"{event.make} {event.model}")
            self.meta_var.set(f"OE: {event.oe or '—'}  ·  Supplier: {event.supplier or '—'}  ·  Freq: {event.freq or '—'}")
            self.row_counter.configure(text=f"ROW {event.excel_row} / {event.total}")
            self.code_vars["A"].set(event.code_a or "—")
            self.code_vars["B"].set(event.code_b or "—")
            self.code_vars["C"].set(event.code_c or "—")
            # Keep last Board telemetry on screen while the next row programs/triggers.
            if getattr(self, "_last_board_tel", None) and hasattr(self, "tel_phase_var"):
                self.tel_phase_var.set("Programming ABC… · showing last Board reading")
            elif hasattr(self, "tel_phase_var"):
                self.tel_phase_var.set("Programming ABC…")
            self._update_stat_labels(event.pending, event.done, event.ok, event.nok, event.total)
            self.status_var.set(f"Testing row {event.excel_row}: {event.make} {event.model} · {event.pending} pending")
            # Board-only USB-TTL: never start/retune rtl_433 (TPMS Board ≠ SDR).
            if self._sdr_enabled() and callable(self.on_prepare_sdr):
                try:
                    from tpms_bench.rtl433 import parse_freq_hz

                    mhz = parse_freq_hz(event.freq) / 1_000_000.0 if event.freq else None
                    self.on_prepare_sdr(freq_mhz=mhz)
                except Exception:
                    pass
            if self._sdr_enabled() and self.on_board_trigger:
                try:
                    self.on_board_trigger(
                        sensor_id=str(event.sensor_id or ""),
                        excel_row=event.excel_row,
                        vehicle=f"{event.make} {event.model}".strip(),
                    )
                except Exception:
                    pass

        elif event.kind == "board_ok":
            # Board ABC finished — always refresh live telemetry tiles from USB-TTL.
            perf = (event.performance or "").upper()
            self._set_telemetry(
                sensor_id=event.sensor_id,
                temperature=str(event.temperature or ""),
                pressure=str(event.pressure or ""),
                rssi=str(event.rssi or ""),
                voltage=event.voltage,
                battery_percentage=event.battery_percentage,
                phase=f"Board {perf or 'read'} · USB-TTL live",
            )
            self._push_tpms_range(event)
            if self._sdr_enabled():
                self.status_var.set(
                    event.message
                    or (
                        f"Row {event.excel_row} Board {perf or '—'} — "
                        "waiting rtl_433 before next ABC…"
                    )
                )
                self._note_sdr_compare("PENDING")
            else:
                self.status_var.set(
                    event.message
                    or f"Row {event.excel_row} Board {perf or '—'} · USB-TTL telemetry updated"
                )
                self._note_sdr_compare("na")
            if self._sdr_enabled() and self.on_board_rf and event.sensor_id and str(event.sensor_id).lower() not in {"na", "n/a"}:
                try:
                    self.on_board_rf(
                        event.sensor_id,
                        excel_row=event.excel_row,
                        vehicle=f"{event.make} {event.model}".strip(),
                        temperature=event.temperature,
                        voltage=event.voltage,
                        pressure=event.pressure,
                        board_result=event.performance or "OK",
                    )
                except Exception:
                    pass
            if self._sdr_enabled() and self.on_board_trigger:
                try:
                    self.on_board_trigger(
                        sensor_id=str(event.sensor_id or ""),
                        excel_row=event.excel_row,
                        vehicle=f"{event.make} {event.model}".strip(),
                    )
                except Exception:
                    pass

        elif event.kind == "row_done":
            self._update_stat_labels(event.pending, event.done, event.ok, event.nok, event.total)
            perf = (event.performance or "").upper()
            self._set_telemetry(
                sensor_id=event.sensor_id,
                temperature=str(event.temperature or ""),
                pressure=str(event.pressure or ""),
                rssi=str(event.rssi or ""),
                voltage=event.voltage,
                battery_percentage=event.battery_percentage,
                phase=f"Row done · Board {perf or '—'}",
            )
            self._push_tpms_range(event)
            vals = self._board_match_values(
                excel_row=int(event.excel_row or 0),
                vehicle=f"{event.make} {event.model}".strip(),
                codes=f"{event.code_a}  {event.code_b}  {event.code_c}".strip(),
                sensor_id=str(event.sensor_id or "na"),
                performance=str(event.performance or "SKIP"),
                temperature=str(event.temperature or "na"),
                voltage=_fmt_batt(event.voltage, event.battery_percentage).removeprefix("Batt ").strip()
                or "na",
                pressure=str(event.pressure or "na"),
                rssi=str(event.rssi or "na"),
                reason=str(event.reason or event.message or ""),
            )
            tags = (event.performance or "SKIP",)
            if not hasattr(self, "_history_all"):
                self._history_all = []
            self._history_all.insert(0, (vals, tags))
            item = self.tree.insert("", 0, values=vals, tags=tags)
            self._flash_row(item, event.performance or "SKIP")
            self.tree.see(item)
            self._record_session_row(event)
            self._note_sdr_compare(event.sdr_compare)
            self._apply_history_filter()
            level = "ok" if (event.performance or "").upper() == "OK" else (
                "error" if (event.performance or "").upper() == "NOK" else "info"
            )
            self._log_board(
                f"Row {event.excel_row}  {event.performance or '—'}  "
                f"ID {event.sensor_id or 'na'}  "
                f"{event.temperature or '—'}°C  "
                f"{event.pressure or '—'}  "
                f"{_fmt_batt(event.voltage, event.battery_percentage)}",
                level=level,
            )
            if hasattr(self, "charts"):
                self.charts.add_result(event.performance or "")
            if hasattr(self, "trend_charts"):
                volt = None
                temp = None
                try:
                    v_raw = str(getattr(event, "voltage", "") or "").strip()
                    if v_raw and v_raw.lower() not in {"na", "n/a", "—", "-"}:
                        volt = float(v_raw)
                except (TypeError, ValueError):
                    volt = None
                try:
                    t_raw = str(getattr(event, "temperature", "") or "").strip()
                    if t_raw and t_raw.lower() not in {"na", "n/a", "—", "-"}:
                        temp = float(t_raw)
                except (TypeError, ValueError):
                    temp = None
                try:
                    self.trend_charts.push(
                        result=event.performance or "",
                        voltage=volt,
                        temp=temp,
                    )
                except Exception:
                    pass
            # Stamp final rtl_433 library decoder only when this session uses SDR.
            finalize_cb = getattr(self, "on_sdr_finalize", None)
            if self._sdr_enabled() and callable(finalize_cb):
                try:
                    finalize_cb(
                        excel_row=event.excel_row,
                        sensor_id=event.sensor_id,
                        decoder=event.rtl433_decoder,
                        sdr_compare=event.sdr_compare,
                        sdr_reason=event.sdr_reason,
                    )
                except Exception:
                    pass
            if self._sdr_enabled():
                sdr = (event.sdr_compare or "").upper()
                self.status_var.set(
                    f"Row {event.excel_row} Board {event.performance or '—'} · SDR {sdr or '—'} "
                    f"— next ABC…"
                )
            else:
                self.status_var.set(
                    f"Row {event.excel_row} Board {event.performance or '—'} · USB-TTL — next ABC…"
                )

        elif event.kind == "stopping":
            self._set_state_badge("STOPPED")
            self.status_var.set(event.message or "Stopping…")

        elif event.kind == "chunk_complete":
            self.is_paused = True
            self.pause_btn.configure(text="Resume")
            self._set_state_badge("PAUSED")
            self.status_var.set(event.message or "Chunk complete — review Comparison")
            self._release_sdr_after_chunk(event.message or "")
            self.after(80, lambda: self.on_chunk_complete(event.message or ""))

        elif event.kind == "finished":
            self._end_session("COMPLETED", event.message)

        elif event.kind == "stopped":
            self._end_session("STOPPED", event.message)

        elif event.kind == "error":
            # If operator already hit Stop, treat COM/open failures as a clean stop.
            if self._state_label == "STOPPED" or (self.runner and self.runner.stop_flag):
                self._end_session("STOPPED", event.message or "Stopped")
            else:
                self.start_btn.configure(state="normal")
                if hasattr(self, "full_btn"):
                    self.full_btn.configure(state="normal")
                self.pause_btn.configure(state="disabled", text="Pause test")
                self.stop_btn.configure(state="disabled")
                self.is_paused = False
                self._set_state_badge("ERROR")
                messagebox.showerror("TPMS Board Error", event.message, parent=self._toplevel())
                self.status_var.set(event.message or "Error")

    def _update_stat_labels(self, pending: int, done: int, ok: int, nok: int, total: int) -> None:
        self.stat_pending.set_value(f"{pending:,}" if isinstance(pending, int) else str(pending))
        self.stat_done.set_value(str(done))
        self.stat_ok.set_value(str(ok))
        self.stat_nok.set_value(str(nok))
        rate_pct = (ok / done * 100.0) if done else 0.0
        self.stat_rate.set_value(f"{rate_pct:.1f}%")
        if hasattr(self.stat_pending, "set_subtitle"):
            self.stat_pending.set_subtitle("in queue")
            self.stat_done.set_subtitle(f"of {total:,} rows" if total else "of 0 rows")
            self.stat_ok.set_subtitle("all codes matched" if nok == 0 and ok > 0 else "OK rows")
            self.stat_nok.set_subtitle("no failures" if nok == 0 else "need review")
            self.stat_rate.set_subtitle("")
        if total > 0:
            pct = done / total
            self.progress.set(min(1.0, pct))
            self.progress_label.configure(text=f"Progress: {pct * 100:.1f}% ({done} of {total})")
            if self.start_time and done > 0:
                elapsed_s = time.monotonic() - self.start_time
                mins, secs = divmod(int(elapsed_s), 60)
                rate = done / (elapsed_s / 60.0) if elapsed_s > 0 else 0
                self.speed_label.configure(text=f"{rate:.1f} tests/min")
                self.elapsed_label.configure(text=f"Elapsed {mins}:{secs:02d}")
                if hasattr(self, "eta_label") and rate > 0 and pending > 0:
                    eta_m = pending / rate
                    if eta_m >= 60:
                        h, m = divmod(int(eta_m), 60)
                        self.eta_label.configure(text=f"ETA ~{h} h {m:02d} m")
                    else:
                        self.eta_label.configure(text=f"ETA ~{int(eta_m)} m")
                elif hasattr(self, "eta_label"):
                    self.eta_label.configure(text="")
        else:
            self.progress.set(0)

    def shutdown(self) -> None:
        if self.runner:
            self.runner.request_stop()
